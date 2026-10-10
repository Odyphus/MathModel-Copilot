"""Declared CSV/XLSX normalization. No guessed units, dates, gaps or imputation.

Only produces local inputs; registration and mathematical verification remain
separate transactions in the existing project authority.
"""
from __future__ import annotations

import argparse
import csv
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from fractions import Fraction
import hashlib
import io
import json
import math
from pathlib import Path
import re

from copilot_domain import seal_record


class DataError(ValueError):
    """An input or its declared interpretation cannot be used safely."""


def project_path(root, reference, *, exists=True):
    if not isinstance(reference, str) or not reference.strip():
        raise DataError("必须提供项目相对路径")
    relative = Path(reference)
    if relative.is_absolute() or ".." in relative.parts or ":" in reference or "\\" in reference:
        raise DataError("路径必须是使用 / 的安全项目相对路径")
    result = (Path(root).resolve() / relative).resolve()
    if not result.is_relative_to(Path(root).resolve()):
        raise DataError("路径越出项目目录")
    if exists and not result.is_file():
        raise DataError("文件不存在：" + reference)
    return result


def binding(root, reference):
    path = project_path(root, reference)
    content = path.read_bytes()
    return {"path": reference, "sha256": hashlib.sha256(content).hexdigest().upper(), "byte_size": len(content)}


def read_spec(root, reference):
    try:
        spec = json.loads(project_path(root, reference).read_text(encoding="utf-8-sig"))
    except (json.JSONDecodeError, UnicodeError) as exc:
        raise DataError("数据解释文件必须是 UTF-8 JSON") from exc
    if not isinstance(spec, dict) or spec.get("schema_version") != "1.0":
        raise DataError("数据解释 schema_version 必须为 1.0")
    def allowed(value, fields, label):
        if not isinstance(value, dict) or set(value) - set(fields):
            raise DataError(label + " 含未知字段或不是对象；检查拼写，不能忽略声明")
    allowed(spec, {"schema_version", "question", "observation_unit", "source", "columns", "unique_keys", "time_axis", "split"}, "数据解释")
    allowed(spec.get("source"), {"path", "sheet", "encoding", "header_row"}, "source")
    if not re.fullmatch(r"Q[1-9]\d*", str(spec.get("question", ""))):
        raise DataError("question 必须为 Q1、Q2 等小问标识")
    columns = spec.get("columns")
    if not isinstance(columns, list) or not columns:
        raise DataError("columns 必须声明非空列清单")
    if any(not isinstance(c, dict) for c in columns):
        raise DataError("每个列声明必须是对象")
    for c in columns:
        allowed(c, {"name", "source", "type", "unit", "output_unit", "min", "max", "nullable", "fill", "format", "meaning", "power_basis"}, "column")
        if not isinstance(c.get("name"), str) or not c["name"].strip() or not isinstance(c.get("source"), str):
            raise DataError("列须声明 name 和 source")
        if c.get("type") not in {"number", "integer", "string", "date", "time"}:
            raise DataError("不支持的列类型：" + c["name"])
        if c.get("fill", "reject") not in {"reject", "forward", "merged"}:
            raise DataError("不支持的填充策略")
        if c.get("fill", "reject") != "reject" and c["type"] != "date":
            raise DataError("前向填充只允许明确声明的日期列")
        if c["type"] in {"number", "integer"} and not isinstance(c.get("unit"), str):
            raise DataError("数值列必须明确 unit；无量纲填写 1")
        if not isinstance(c.get("nullable", False), bool):
            raise DataError("nullable 必须为布尔值")
        for key in ("min", "max"):
            if c.get(key) is not None and (isinstance(c[key], bool) or not isinstance(c[key], (int, float)) or not math.isfinite(c[key])):
                raise DataError("min/max 须为归一化单位下的有限数值")
        if c.get("min") is not None and c.get("max") is not None and c["min"] > c["max"]:
            raise DataError("min 不得大于 max")
    names = [c["name"] for c in columns]
    sources = [c["source"] for c in columns]
    if len(set(names)) != len(names) or len(set(sources)) != len(sources):
        raise DataError("列名称或源列重复")
    if set(names) & {"slot_start", "slot_end", "source_row"}:
        raise DataError("列名占用了适配器保留字段")
    keys = spec.get("unique_keys", [])
    if not isinstance(keys, list) or any(k not in names for k in keys):
        raise DataError("unique_keys 必须引用已声明列")
    axis = spec.get("time_axis")
    if axis is not None:
        allowed(axis, {"date_column", "time_column", "label", "interval_minutes", "start", "periods"}, "time_axis")
        if not isinstance(axis, dict) or axis.get("label") not in {"start", "end"}:
            raise DataError("time_axis 须明确时标是 start 还是 end")
        types = {c["name"]: c["type"] for c in columns}
        if types.get(axis.get("date_column")) != "date" or types.get(axis.get("time_column")) != "time":
            raise DataError("time_axis 须引用日期列和时间列")
        interval = axis.get("interval_minutes")
        periods = axis.get("periods")
        if isinstance(interval, bool) or not isinstance(interval, (int, float)) or not math.isfinite(interval) or interval <= 0:
            raise DataError("interval_minutes 必须是有限正数")
        if isinstance(periods, bool) or not isinstance(periods, int) or periods <= 0:
            raise DataError("periods 必须是正整数")
        start = _datetime(axis.get("start"))
        if start.tzinfo is not None:
            raise DataError("本适配器只接受无时区的本地日历时间；时区转换须另行声明处理")
    return spec


def _datetime(value):
    try:
        return datetime.fromisoformat(value)
    except (ValueError, TypeError) as exc:
        raise DataError("时间起点须使用 ISO 日期时间") from exc


def _read_table(root, source):
    if not isinstance(source, dict):
        raise DataError("source 必须是对象")
    path = project_path(root, source.get("path"))
    header_row = source.get("header_row", 1)
    if isinstance(header_row, bool) or not isinstance(header_row, int) or header_row < 1:
        raise DataError("header_row 必须是正整数")
    if path.suffix.lower() == ".csv":
        try:
            with path.open(encoding=source.get("encoding", "utf-8-sig"), newline="") as handle:
                rows = list(csv.reader(handle))
        except (UnicodeError, LookupError) as exc:
            raise DataError("CSV 编码不匹配；请明确 source.encoding") from exc
        if len(rows) < header_row:
            raise DataError("表格没有声明的表头行")
        return rows[header_row - 1], [(n, row) for n, row in enumerate(rows[header_row:], header_row + 1)], {}, set(), None
    if path.suffix.lower() != ".xlsx":
        raise DataError("仅支持 .csv / .xlsx；旧版 XLS 和 PDF 尚未包含在此适配器中")
    try:
        import openpyxl
    except ImportError as exc:
        raise DataError("读取 XLSX 需要 openpyxl；安装 analysis extra 或 openpyxl>=3.1") from exc
    values = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)
    try:
        sheet_name = source.get("sheet")
        if not isinstance(sheet_name, str) or sheet_name not in values.sheetnames:
            raise DataError("XLSX 必须明确存在的 source.sheet，不能猜第一张表")
        sheet = values[sheet_name]
        formula_sheet = formulas[sheet_name]
        header = [cell.value for cell in sheet[header_row]]
        rows = [(n, [cell.value for cell in sheet[n]]) for n in range(header_row + 1, sheet.max_row + 1)]
        merge_anchors = {}
        for area in formula_sheet.merged_cells.ranges:
            for n in range(max(header_row + 1, area.min_row), area.max_row + 1):
                for col in range(area.min_col, area.max_col + 1):
                    merge_anchors[n, col] = (area.min_row, area.min_col)
        missing_formula_cache = {(cell.row, cell.column) for row in formula_sheet for cell in row
                                 if cell.data_type == "f" and sheet.cell(cell.row, cell.column).value is None}
        return header, rows, merge_anchors, missing_formula_cache, sheet_name
    finally:
        values.close()
        formulas.close()


def _is_missing(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _date(value, column):
    if isinstance(value, datetime):
        if value.tzinfo is not None or value.time() != time(0):
            raise DataError("日期列包含时刻/时区，不能悄悄丢弃")
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if not isinstance(value, str):
        raise DataError("日期须为 Excel 日期类型或声明格式的文字")
    try:
        return datetime.strptime(value.strip(), column.get("format", "%Y-%m-%d")).date().isoformat()
    except ValueError as exc:
        raise DataError("日期不符合声明格式") from exc


def _time(value):
    if isinstance(value, time):
        if value.tzinfo is not None:
            raise DataError("时间列不支持时区")
        return value.isoformat(timespec="microseconds" if value.microsecond else "seconds")
    if not isinstance(value, str):
        raise DataError("时间须为 Excel time 或 HH:MM[:SS][+1] 文字；不猜数值序列号")
    match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::(\d{2}))?(\+1)?", value.strip())
    if not match:
        raise DataError("时间格式应为 HH:MM[:SS][+1]")
    hour, minute, second = (int(match.group(i) or 0) for i in (1, 2, 3))
    offset = match.group(4) or ""
    if hour == 24 and minute == second == 0 and not offset:
        hour, offset = 0, "+1"
    try:
        return time(hour, minute, second).isoformat() + offset
    except ValueError as exc:
        raise DataError("时间超出有效范围") from exc


def _unit_factor(column, interval):
    source = column["unit"].strip()
    target = column.get("output_unit", source)
    if not source or not isinstance(target, str) or not target.strip():
        raise DataError("单位不能留空；无量纲填写 1")
    power = {"W": 1, "kW": 1000, "MW": 1000000}
    energy = {"Wh": 1, "kWh": 1000, "MWh": 1000000}
    if source == target:
        return 1.0
    for dimension in (power, energy):
        if source in dimension and target in dimension:
            return dimension[source] / dimension[target]
    if source in power and target in energy and interval is not None:
        if column.get("power_basis") != "interval_mean":
            raise DataError("功率转区间能量须明确 power_basis=interval_mean；瞬时测量不能直接视为区间平均")
        return power[source] * interval / 60 / energy[target]
    raise DataError(f"不支持或缺少区间时长的单位换算：{source} → {target}")


def _parse(value, column, interval):
    if _is_missing(value):
        if column.get("nullable", False) and column["type"] not in {"date", "time"}:
            return None
        raise DataError("缺失值未获准保留；不能补零")
    kind = column["type"]
    if kind == "date":
        return _date(value, column)
    if kind == "time":
        return _time(value)
    if kind == "string":
        if not isinstance(value, str):
            raise DataError("声明的文字列不是文字")
        return value
    if isinstance(value, bool) or isinstance(value, (date, time)):
        raise DataError("数值列不能使用布尔值或日期时间")
    if kind == "integer":
        try:
            exact = Decimal(str(value))
            if not exact.is_finite() or exact != exact.to_integral_value():
                raise DataError("声明的整数列包含非整数或非有限值")
            scaled = Fraction(int(exact)) * Fraction(str(_unit_factor(column, interval)))
            if scaled.denominator != 1:
                raise DataError("单位换算后的整数列不再是整数；应改为 number")
            number = scaled.numerator
        except InvalidOperation as exc:
            raise DataError("无法解释为整数") from exc
        if column.get("min") is not None and number < column["min"] or column.get("max") is not None and number > column["max"]:
            raise DataError("整数超出归一化单位下声明的范围")
        return number
    try:
        number = float(value)
    except (ValueError, TypeError) as exc:
        raise DataError("无法解释为数值") from exc
    if not math.isfinite(number):
        raise DataError("数值必须有限，不能使用 NaN 或 Inf")
    number *= _unit_factor(column, interval)
    if not math.isfinite(number):
        raise DataError("单位换算后数值不是有限值")
    if column.get("min") is not None and number < column["min"]:
        raise DataError("数值低于归一化单位下的 min")
    if column.get("max") is not None and number > column["max"]:
        raise DataError("数值高于归一化单位下的 max")
    return number


def normalize(root, spec):
    header, source_rows, merges, formula_missing, sheet = _read_table(root, spec.get("source"))
    if any(not isinstance(h, str) or not h.strip() for h in header) or len(set(header)) != len(header):
        raise DataError("表头为空或重复，须先明确附件的列结构")
    positions = {c["name"]: header.index(c["source"]) for c in spec["columns"] if c["source"] in header}
    if len(positions) != len(spec["columns"]):
        raise DataError("附件缺少声明的源列")
    axis = spec.get("time_axis")
    interval = axis["interval_minutes"] if axis else None
    rows, locations, fills, previous = [], [], [], {}
    seen_keys, seen_slots = set(), set()
    for number, cells in source_rows:
        if len(cells) != len(header):
            raise DataError(f"第 {number} 行列数与表头不同")
        row, origins = {}, {}
        for column in spec["columns"]:
            name = column["name"]
            index = positions[name]
            origin = (number, index + 1)
            value = cells[index]
            if origin in formula_missing:
                raise DataError(f"第 {number} 行 {column['source']} 公式缺少缓存，须先由电子表格软件重算")
            if _is_missing(value) and column.get("fill") == "merged" and origin in merges:
                anchor = merges[origin]
                if anchor[1] != index + 1 or name not in previous or previous[name][1] != anchor:
                    raise DataError(f"第 {number} 行的合并日期不是本列已读取的锚点")
                value, origin = previous[name]
            elif _is_missing(value) and column.get("fill") == "forward" and name in previous:
                value, origin = previous[name]
            if origin != (number, index + 1):
                fills.append({"row": number, "column": name, "from_row": origin[0], "policy": column["fill"]})
            try:
                row[name] = _parse(value, column, interval)
            except DataError as exc:
                raise DataError(f"第 {number} 行 {column['source']}：{exc}") from exc
            if not _is_missing(value):
                previous[name] = (value, origin)
            origins[name] = {"row": origin[0], "column": origin[1]}
        if spec.get("unique_keys"):
            key = tuple(row[k] for k in spec["unique_keys"])
            if any(x is None for x in key) or key in seen_keys:
                raise DataError(f"第 {number} 行缺少唯一键或唯一键重复")
            seen_keys.add(key)
        if axis:
            clock = row[axis["time_column"]]
            offset = 1 if clock.endswith("+1") else 0
            stamp = datetime.combine(date.fromisoformat(row[axis["date_column"]]), time.fromisoformat(clock.removesuffix("+1"))) + timedelta(days=offset)
            step = timedelta(minutes=interval)
            start = stamp if axis["label"] == "start" else stamp - step
            if start in seen_slots:
                raise DataError(f"第 {number} 行时间槽重复")
            expected = _datetime(axis["start"]) + len(rows) * step
            if start != expected:
                raise DataError(f"第 {number} 行时间槽缺失、乱序或偏移：应为 {expected.isoformat()}，实际 {start.isoformat()}")
            seen_slots.add(start)
            row.update(slot_start=start.isoformat(), slot_end=(start + step).isoformat())
        row["source_row"] = number
        rows.append(row)
        locations.append({"source_row": number, "columns": origins})
    if not rows:
        raise DataError("没有数据行")
    if axis and len(rows) != axis["periods"]:
        raise DataError(f"时间槽数量不符：预期 {axis['periods']}，实际 {len(rows)}")
    fields = [{"name": c["name"], "source": c["source"], "data_type": c["type"],
               "meaning": c.get("meaning", c["source"]), "unit": c.get("output_unit", c.get("unit", "calendar" if c["type"] in {"date", "time"} else "text")),
               "source_unit": c.get("unit"), "factor": _unit_factor(c, interval) if c["type"] in {"number", "integer"} else None,
               "power_basis": c.get("power_basis"),
               "nullable": c.get("nullable", False), "fill": c.get("fill", "reject")}
              for c in spec["columns"]]
    return rows, {"schema_version": "1.0", "row_count": len(rows), "sheet": sheet, "fields": fields,
                  "time_axis": axis, "filled_dates": fills, "locations": locations,
                  "scope": "Declared parsing, units, keys and slot alignment only; no model or scientific validation"}


def csv_bytes(rows):
    handle = io.StringIO(newline="")
    writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return handle.getvalue().encode("utf-8")


def verify_csv(path, rows):
    # Reopen the actual file, checking values/slots rather than only existence.
    with Path(path).open(encoding="utf-8", newline="") as handle:
        actual = list(csv.DictReader(handle))
    if len(actual) != len(rows) or list(actual[0]) != list(rows[0]):
        raise DataError("写后读回的行数或列映射不一致")
    for number, (written, expected) in enumerate(zip(actual, rows), 1):
        for key, value in expected.items():
            if written[key] != ("" if value is None else str(value)):
                raise DataError(f"写后读回第 {number} 槽 {key} 与归一化解不一致")


def prepare(root, spec_reference, output):
    root = Path(root).resolve()
    spec = read_spec(root, spec_reference)
    source_binding = binding(root, spec["source"]["path"])
    spec_binding = binding(root, spec_reference)
    rows, report = normalize(root, spec)
    target = project_path(root, output, exists=False)
    if target.exists():
        raise DataError("输出目录已存在；为新版本指定新目录，不覆盖原文件")
    target.mkdir(parents=True)
    data_ref = output.rstrip("/") + "/normalized.csv"
    report_ref = output.rstrip("/") + "/report.json"
    contract_ref = output.rstrip("/") + "/data_contract.json"
    project_path(root, data_ref, exists=False).write_bytes(csv_bytes(rows))
    verify_csv(root / data_ref, rows)
    if binding(root, spec["source"]["path"]) != source_binding or binding(root, spec_reference) != spec_binding:
        raise DataError("读入期间源文件或解释文件发生变化，请重试新版本")
    report.update(source=source_binding, spec=spec_binding, normalized=binding(root, data_ref), readback={"passed": True, "criterion": "All slots and all serialized values equal normalized records", "tolerance": 0, "unit": "serialized cell"})
    (root / report_ref).write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    contract = {"question": spec["question"],
                "inventory": seal_record({"entries": [source_binding, spec_binding, binding(root, data_ref), binding(root, report_ref)]}),
                "passport": seal_record({"fields": report["fields"], "observation_unit": spec.get("observation_unit", "declared table row or time interval"),
                                         "missing_policy": "reject unless explicitly nullable; declared date fill only", "outlier_policy": "reject declared bound violations; preserve all other values"}),
                "split": seal_record(spec.get("split", {"applicability": "not_applicable", "not_applicable_reason": "Normalized input for deterministic computation only; no learned or predictive claim",
                                                        "evaluation_target": "deterministic_computation", "claims_new_entity_generalization": False, "claims_predictive_performance": False})),
                "adapter": {"schema_version": "1.0", "spec": spec_binding, "report": binding(root, report_ref), "normalized": binding(root, data_ref)}}
    (root / contract_ref).write_text(json.dumps(contract, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    return {"row_count": len(rows), "contract": contract_ref, "normalized": data_ref, "report": report_ref, "registered": False, "readback_passed": True}


def validate_adapter(root, contract):
    """Replay declared normalization before accepting its optional contract binding."""
    adapter = contract.get("adapter")
    if adapter is None:
        return []
    try:
        if not isinstance(adapter, dict) or adapter.get("schema_version") != "1.0":
            raise DataError("adapter schema_version 必须为 1.0")
        bound = []
        for key in ("spec", "report", "normalized"):
            item = adapter[key]
            if not isinstance(item, dict) or binding(root, item["path"]) != item:
                raise DataError("适配器文件绑定不一致：" + key)
            bound.append(item)
        spec = read_spec(root, adapter["spec"]["path"])
        if spec["question"] != contract.get("question"):
            raise DataError("数据解释与 DataContract 的小问不一致")
        rows, expected = normalize(root, spec)
        if (Path(root) / adapter["normalized"]["path"]).read_bytes() != csv_bytes(rows):
            raise DataError("归一化文件与原始附件按声明转换的结果不一致")
        verify_csv(Path(root) / adapter["normalized"]["path"], rows)
        report = json.loads(project_path(root, adapter["report"]["path"]).read_text(encoding="utf-8"))
        expected.update(source=binding(root, spec["source"]["path"]), spec=adapter["spec"], normalized=adapter["normalized"],
                        readback={"passed": True, "criterion": "All slots and all serialized values equal normalized records", "tolerance": 0, "unit": "serialized cell"})
        if report != expected or contract["passport"]["fields"] != expected["fields"]:
            raise DataError("数据报告或字段护照与本次重读不一致")
        inventory = contract["inventory"]["entries"]
        if any(item not in inventory for item in bound + [expected["source"]]):
            raise DataError("原始附件、解释、报告及归一化文件必须全部绑定到 inventory")
    except (DataError, OSError, KeyError, TypeError, ValueError) as exc:
        return ["DataContract adapter: " + str(exc)]
    return []


def parser():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--workspace", type=Path, default=Path.cwd())
    sub = p.add_subparsers(dest="action", required=True)
    for action in ("inspect", "prepare"):
        s = sub.add_parser(action)
        s.add_argument("--spec", required=True, help="项目内 UTF-8 JSON；示例 templates/copilot/table_spec.json")
        if action == "prepare":
            s.add_argument("--output", required=True, help="项目内不存在的输出目录；不覆盖、不登记或核验数学结果")
    return p


def execute(args):
    if args.action == "prepare":
        return prepare(args.workspace, args.spec, args.output)
    rows, report = normalize(args.workspace, read_spec(args.workspace, args.spec))
    report["filled_date_count"] = len(report["filled_dates"])
    report["locations_truncated"] = len(report["locations"]) > 5
    report["filled_dates_truncated"] = len(report["filled_dates"]) > 5
    report["locations"] = report["locations"][:5]
    report["filled_dates"] = report["filled_dates"][:5]
    return {"report": report, "preview": rows[:5], "written": False}


if __name__ == "__main__":
    import sys
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    try:
        print(json.dumps({"ok": True, "result": execute(parser().parse_args())}, ensure_ascii=False, indent=2, allow_nan=False))
    except (DataError, OSError, KeyError, TypeError, ValueError) as exc:
        print(json.dumps({"ok": False, "message": str(exc)}, ensure_ascii=False))
        raise SystemExit(2)
