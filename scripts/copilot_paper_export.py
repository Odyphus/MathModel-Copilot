"""Export existing verified sections to native Word, then audit the readback.

No prose is generated or rewritten. Optional PDF uses local LibreOffice or an
idle Windows Word COM instance, never a cloud conversion service. Formatting
and conversion checks do not establish contest compliance or human finality.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import unicodedata

from copilot_paper_source import CLAIM, IMAGE, _inline, _table_cells, audit_paper_source, paper_source_blocks
from copilot_runtime import Runtime, bind_file, safe_path
from copilot_section import section_projection


def native_math(expression):
    from docx.oxml import OxmlElement
    math = OxmlElement("m:oMath")
    def run(text):
        node = OxmlElement("m:r"); value = OxmlElement("m:t"); value.text = text; node.append(value); return node
    fraction = re.fullmatch(r"([^\\{}_^]+?)=\\frac\{([^\\{}_^]+)\}\{([^\\{}_^]+)\}", expression)
    if fraction:
        math.append(run(fraction[1] + "=")); part = OxmlElement("m:f")
        for tag, text in (("m:num", fraction[2]), ("m:den", fraction[3])):
            child = OxmlElement(tag); child.append(run(text)); part.append(child)
        math.append(part)
    elif re.fullmatch(r"[^\\{}_^]+", expression):
        math.append(run(expression))
    else:
        raise ValueError("原生公式导出目前仅支持无上下标的表达式及简单分式；复杂公式须使用已验收的专用排版路径")
    return math


def table_layout(rows, available_points):
    """Keep numeric cells on one line without changing their precision.

    Reserve width for body values, allow prose and headers to wrap, and reject
    tables that cannot fit at a readable font size instead of clipping them.
    The estimate is conservative for the declared Calibri/Chinese font pair;
    actual page rendering remains a separate acceptance step.
    """
    numeric = re.compile(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?%?")
    def points(value, size):
        return sum(size if unicodedata.east_asian_width(ch) in ("W", "F")
                   else size * .6 for ch in value)
    for size in (11, 10.5, 10, 9.5, 9):
        widths = []
        for column in zip(*rows):
            header = min(points(column[0], size), size * 4)
            values = [points(value, size) if numeric.fullmatch(value)
                      else min(points(value, size), size * 12) for value in column[1:]]
            widths.append(max(42, max([header, *values]) + 10))
        if sum(widths) <= available_points:
            remaining = (available_points - sum(widths)) / len(widths)
            return [width + remaining for width in widths], size
    raise ValueError("表格列数或数字过长，当前纵向页面无法清晰容纳；请显式调整表结构或选择其他版式，不会自动删减数值精度")


class Renderer:
    def __init__(self, root):
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        self.root = root; self.doc = Document(); self.first = True
        sec = self.doc.sections[0]
        sec.page_width = Inches(8.5); sec.page_height = Inches(11)
        sec.top_margin = sec.bottom_margin = Inches(.8)
        sec.left_margin = sec.right_margin = Inches(.85)
        for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
            style = self.doc.styles[name]; style.font.name = "Calibri"; style.font.color.rgb = RGBColor(0, 0, 0)
            fonts = style._element.get_or_add_rPr().find(qn("w:rFonts"))
            if fonts is None: fonts = OxmlElement("w:rFonts"); style._element.get_or_add_rPr().append(fonts)
            fonts.set(qn("w:eastAsia"), "宋体")
        self.doc.styles["Normal"].font.size = Pt(11)
        self.doc.styles["Normal"].paragraph_format.line_spacing = 1.3
        self.doc.styles["Normal"].paragraph_format.space_after = Pt(7)
        for name, size in (("Title", 19), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
            style = self.doc.styles[name]; style.font.size = Pt(size); style.font.bold = True
            style.paragraph_format.space_before = Pt(12); style.paragraph_format.space_after = Pt(6)
        for border in list(self.doc.styles.element.iter(qn("w:pBdr"))): border.getparent().remove(border)

    def table(self, rows):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        table = self.doc.add_table(rows=0, cols=len(rows[0])); table.autofit = False
        section = self.doc.sections[0]
        widths, font_size = table_layout(rows, (section.page_width-section.left_margin-section.right_margin)/12700)
        for column, width in zip(table.columns, widths): column.width = Pt(width)
        borders = OxmlElement("w:tblBorders")
        for tag in ("top", "bottom", "left", "right", "insideH", "insideV"):
            node = OxmlElement("w:"+tag)
            for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")): node.set(qn("w:"+key), value)
            borders.append(node)
        table._tbl.tblPr.append(borders)
        for i, values in enumerate(rows):
            row = table.add_row(); row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if i == 0: row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for j, (cell, value) in enumerate(zip(row.cells, values)):
                cell.width = Pt(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                properties = cell._tc.get_or_add_tcPr()
                if i == 0:
                    shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "E7EBEE"); properties.append(shade)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "bottom", "left", "right"):
                    node = OxmlElement("w:"+side); node.set(qn("w:w"), "80" if side in ("left", "right") else "110"); node.set(qn("w:type"), "dxa"); margins.append(node)
                properties.append(margins)
                p = cell.paragraphs[0]; p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(value); run.bold = i == 0; run.font.size = Pt(font_size)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def markdown(self, text):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches
        lines = CLAIM.sub("", text).splitlines(); index = 0
        while index < len(lines):
            line = lines[index].strip()
            if not line: index += 1; continue
            if line.startswith("|"):
                rows = [_table_cells(line)]; index += 2
                while index < len(lines) and lines[index].strip().startswith("|"):
                    rows.append(_table_cells(lines[index])); index += 1
                self.table(rows); continue
            if line.startswith("$$"):
                match = re.fullmatch(r"\$\$(.+?)\$\$\s*(.*)", line)
                p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p._p.append(native_math(match[1])); p.add_run(_inline(match[2])); index += 1; continue
            image = IMAGE.fullmatch(line)
            if image:
                p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                picture = p.add_run().add_picture(str(safe_path(self.root, image[2])), width=Inches(6.2))
                picture._inline.docPr.set("descr", image[1]); index += 1; continue
            if line.startswith("#"):
                level = min(len(line)-len(line.lstrip("#")), 3)
                p = self.doc.add_paragraph(_inline(re.sub(r"^#{1,6}\s+", "", line)), style="Title" if self.first else "Heading "+str(level))
                if self.first: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self.first = False; index += 1; continue
            parts = []
            while index < len(lines) and lines[index].strip():
                current = lines[index].strip()
                if parts and (current.startswith(("|", "$$", "#")) or IMAGE.fullmatch(current)): break
                parts.append(current); index += 1
            self.doc.add_paragraph(_inline("\n".join(parts)))


def convert_pdf(docx, output):
    """Bounded local conversion; a PDF file is not a visual acceptance receipt."""
    if output.exists(): raise ValueError("PDF 已存在；不覆盖")
    output.parent.mkdir(parents=True, exist_ok=True)
    office = shutil.which("soffice") or shutil.which("soffice.exe")
    if office:
        with tempfile.TemporaryDirectory(prefix="copilot-pdf-") as directory:
            base = Path(directory)
            completed = subprocess.run([office, "-env:UserInstallation="+(base/"profile").as_uri(), "--headless", "--convert-to", "pdf", "--outdir", str(base), str(docx)], capture_output=True, timeout=90)
            made = base / (docx.stem + ".pdf")
            if completed.returncode != 0 or not made.is_file(): raise ValueError("本地 LibreOffice 转换失败")
            with output.open("xb") as handle: handle.write(made.read_bytes())
        backend = "local LibreOffice"
    elif os.name == "nt":
        shell = shutil.which("powershell.exe")
        if not shell: raise ValueError("PDF 转换缺少本地 LibreOffice 或 Windows Word 后端")
        script = Path(__file__).with_name("copilot_word_pdf.ps1")
        completed = subprocess.run([shell, "-NoProfile", "-NonInteractive", "-File", str(script), "-InputDocx", str(docx), "-OutputPdf", str(output)], capture_output=True, timeout=90)
        if completed.returncode != 0 or not output.is_file():
            details = (completed.stderr or completed.stdout).decode("utf-8", errors="replace").strip()[:1600]
            raise ValueError("本地 Word 转换失败或 Word 正在使用；DOCX 已保留：" + details)
        backend = "local Microsoft Word COM"
    else:
        raise ValueError("PDF 转换需要本地 LibreOffice；DOCX 已保留")
    if not output.read_bytes().startswith(b"%PDF-"): raise ValueError("转换产物不是有效 PDF 头")
    return backend


def export(root, contract_ref, output, *, expected_revision, pdf=None):
    root = Path(root).resolve(); rt = Runtime(root); state = rt.read(); cp = state["copilot"]
    if cp["revision"] != expected_revision:
        raise ValueError("项目版本已变化；请读取当前版本后再导出")
    contract_path = safe_path(root, contract_ref)
    contract = json.loads(contract_path.read_text(encoding="utf-8")); sections = contract.get("sections")
    sources = paper_source_blocks(root, cp, sections, contract)
    target = safe_path(root, output, exists=False)
    pdf_path = safe_path(root, pdf, exists=False) if pdf is not None else None
    if pdf_path is not None and (pdf_path.suffix.lower() != ".pdf" or pdf_path.exists()):
        raise ValueError("PDF 须使用未占用的新 .pdf 路径")
    report_ref = output + ".source-audit.json"; report_path = safe_path(root, report_ref, exists=False)
    if target.suffix.lower() != ".docx" or target.exists() or report_path.exists(): raise ValueError("DOCX/报告须使用未占用的新路径")
    before = bind_file(root, "state/decision_log.json")
    contract_binding = bind_file(root, contract_ref)
    renderer = Renderer(root)
    for oid in contract["sections"]:
        payload = cp["objects"][oid]["payload"]
        renderer.markdown(section_projection(cp, safe_path(root, payload["path"]).read_text(encoding="utf-8"), payload.get("source_bindings"), payload.get("structure"), root=root))
    for item in contract.get("supplements", []):
        if item["kind"] == "file_list":
            for path in item["paths"]: renderer.doc.add_paragraph("文件："+path)
        else:
            renderer.doc.add_paragraph("源程序："+item["path"])
            renderer.doc.add_paragraph(safe_path(root, item["path"]).read_text(encoding="utf-8-sig").replace("\r\n", "\n").rstrip("\n"))
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="paper-export-", dir=target.parent) as directory:
        draft = Path(directory) / "draft.docx"; renderer.doc.save(draft)
        audit = audit_paper_source(root, cp, sections, draft, contract)
        if not audit["passed"]: raise ValueError("DOCX 写后读回不一致："+str(audit["errors"]))
        if before != bind_file(root, "state/decision_log.json") or contract_binding != bind_file(root, contract_ref) or any(bind_file(root, f["path"]) != f for f in sources["files"]):
            raise ValueError("论文来源在导出期间变化；未输出混合版本")
        with target.open("xb") as handle: handle.write(draft.read_bytes())
    # Store registers sources so same-revision drift and later edits invalidate it.
    audit["docx_path"] = output; audit["scope"] = "Native DOCX source readback only; rendering, competition compliance and human final review are separate."
    with report_path.open("x", encoding="utf-8") as handle: json.dump(audit, handle, ensure_ascii=False, indent=2)
    record_key = "paper-export." + hashlib.sha256(output.encode("utf-8")).hexdigest()[:24]
    record = rt.register(cp["revision"], "ArtifactRecord", record_key,
                         {"artifact_type": "paper_export", "path": output, "source_contract": contract_ref, "source_readback_passed": True,
                          "visual_review": "not_performed", "scope": audit["scope"]},
                         dependencies=sources["dependencies"], files=[output, report_ref, contract_ref])["result"]["object_id"]
    result = {"docx": output, "source_audit": report_ref, "artifact_id": record, "source_readback_passed": True,
              "visual_review": "not_performed", "submission_ready": False}
    if pdf is not None:
        backend = convert_pdf(target, pdf_path)
        revision = rt.read()["copilot"]["revision"]
        pdf_key = "paper-export." + hashlib.sha256(pdf.encode("utf-8")).hexdigest()[:24]
        pdf_record = rt.register(revision, "ArtifactRecord", pdf_key, {"artifact_type": "paper_export", "path": pdf,
                                 "conversion_backend": backend, "visual_review": "not_performed", "scope": "Local PDF conversion; content/layout not automatically certified"},
                                 dependencies=[record], files=[pdf])["result"]["object_id"]
        result.update(pdf=pdf, pdf_artifact_id=pdf_record, conversion_backend=backend)
    return result


def parser():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--project-root", type=Path, required=True)
    p.add_argument("--expected-revision", type=int, required=True, help="导出所依据的当前项目版本；版本变化时停止")
    p.add_argument("--contract", required=True, help="现有 paper source 0.1 合同：sections 顺序及可选绑定 supplements")
    p.add_argument("--output", required=True, help="新建 .docx 项目相对路径")
    p.add_argument("--pdf", help="可选新建 .pdf；需要本地转换后端，失败仍保留 Word 文件")
    return p


if __name__ == "__main__":
    a = parser().parse_args()
    print(json.dumps(export(a.project_root, a.contract, a.output, expected_revision=a.expected_revision, pdf=a.pdf), ensure_ascii=False, indent=2))
