"""Author native Word documents only from registered current paper sources."""
from __future__ import annotations
import hashlib, json, re, sys
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from paper_paths import HERE, PRODUCT, PROJECT
sys.path.insert(0,str(PRODUCT/'scripts'))
from copilot_runtime import Runtime
from copilot_section import section_projection
from copilot_paper_source import audit_paper_source, paper_source_blocks, CLAIM

def save_json(path,data):
    Path(path).write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def element(tag,**attrs):
    node=OxmlElement(tag)
    for key,value in attrs.items(): node.set(qn('w:'+key),str(value))
    return node
def make_document():
    doc=Document(); sec=doc.sections[0]
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Mm(25.1)
    sec.header_distance=Mm(10);sec.footer_distance=Mm(12)
    sec._sectPr.append(element('w:pgNumType',start=1))
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3'):
        style=doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
        style._element.get_or_add_rPr().append(element('w:rFonts',ascii='Times New Roman',hAnsi='Times New Roman',eastAsia='宋体'))
    normal=doc.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.35;normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',18),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
        style=doc.styles[name];style.font.size=Pt(size);style.font.bold=True
        style.paragraph_format.space_before=Pt(12);style.paragraph_format.space_after=Pt(6)
        style.paragraph_format.keep_with_next=True
    code=doc.styles.add_style('CUMCMSourceCode',WD_STYLE_TYPE.PARAGRAPH)
    code.font.name='Consolas';code.font.size=Pt(8)
    code._element.get_or_add_rPr().append(element('w:rFonts',ascii='Consolas',hAnsi='Consolas',eastAsia='宋体'))
    code.paragraph_format.line_spacing=1.0;code.paragraph_format.space_after=Pt(6)
    code.paragraph_format.widow_control=False
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=element('w:fldSimple',instr=' PAGE ');footer._p.append(field)
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.core_properties.title='智能车间 RGV 调度的历史复现与核验报告'
    doc.core_properties.subject='历史复现与验证用途，非正式参赛稿'
    return doc

class Renderer:
    def __init__(self,doc,registry=None,table_widths=None,table_font_size=9.5):
        self.doc=doc;self.registry=registry;self.table_widths=table_widths;self.table_font_size=table_font_size;self.bookmark_id=0;self.seen_citations=set();self.first=True
    def bookmark(self,p,text,name):
        self.bookmark_id+=1
        start=element('w:bookmarkStart',id=self.bookmark_id,name=name);end=element('w:bookmarkEnd',id=self.bookmark_id)
        p._p.append(start);p.add_run(text);p._p.append(end)
    def text(self,p,text):
        if not self.registry:
            p.add_run(text);return
        bib=re.match(r'^\[(\d+)\] ',text)
        if bib:
            entry=next(r for r in self.registry['references'] if r['number']==int(bib[1]))
            self.bookmark(p,text,entry['bibliography_anchor']);return
        for part in re.split(r'(\[[12]\])',text):
            if re.fullmatch(r'\[[12]\]',part) and part not in self.seen_citations:
                n=int(part[1:-1]);entry=next(r for r in self.registry['citations'] if r['number']==n)
                self.bookmark(p,part,entry['anchor']);self.seen_citations.add(part)
            else:p.add_run(part)
    def add_table(self,rows):
        count=len(rows[0]);table=self.doc.add_table(rows=0,cols=count)
        table.autofit=False
        width=9050
        widths=(list(self.table_widths) if self.table_widths and len(self.table_widths)==count else [1250,1200,1200,850,850,3700] if count==6 else [1900]+[(width-1900)//(count-1)]*(count-1))
        widths[-1]+=width-sum(widths)
        table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:type'),'dxa');table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:w'),str(width))
        table._tbl.tblPr.append(element('w:tblInd',w=0,type='dxa'))
        for grid,w in zip(table._tbl.tblGrid,widths):grid.set(qn('w:w'),str(w))
        borders=element('w:tblBorders')
        for side in ('top','bottom','insideH'):
            borders.append(element('w:'+side,val='single',sz=4,color='B7B7B7'))
        for side in ('left','right','insideV'):borders.append(element('w:'+side,val='nil'))
        table._tbl.tblPr.append(borders)
        for i,values in enumerate(rows):
            row=table.add_row();row._tr.get_or_add_trPr().append(element('w:cantSplit'))
            if i==0:row._tr.get_or_add_trPr().append(element('w:tblHeader'))
            for j,(cell,value) in enumerate(zip(row.cells,values)):
                cell._tc.get_or_add_tcPr().find(qn('w:tcW')).set(qn('w:w'),str(widths[j]))
                cell._tc.get_or_add_tcPr().append(element('w:vAlign',val='center'))
                if i==0:cell._tc.get_or_add_tcPr().append(element('w:shd',fill='E7EBEE'))
                p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                p.paragraph_format.line_spacing=1.1
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                run=p.add_run(value);run.font.size=Pt(self.table_font_size);run.bold=(i==0)
    def markdown(self,source):
        lines=CLAIM.sub('',source).splitlines();i=0
        while i<len(lines):
            line=lines[i].strip()
            if not line:i+=1;continue
            if line.startswith('|'):
                rows=[[x.strip() for x in line[1:-1].split('|')]];i+=2
                while i<len(lines) and lines[i].strip().startswith('|'):
                    rows.append([x.strip() for x in lines[i].strip()[1:-1].split('|')]);i+=1
                self.add_table(rows);continue
            image=re.fullmatch(r'!\[([^\]]*)\]\(([^\s)]+)\)',line)
            if image:
                p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                inline=p.add_run().add_picture(str(PROJECT/image[2]),width=Mm(154))
                inline._inline.docPr.set('descr',image[1]);i+=1;continue
            if line.startswith('#'):
                title=re.sub(r'^#{1,6}\s+','',line)
                if self.first:
                    p=self.doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;self.first=False
                else:
                    level=min(len(line)-len(line.lstrip('#')),3)
                    p=self.doc.add_paragraph(style='Heading '+str(level))
                    if title.startswith(('1 问题重述','4 实验结果')) or title=='附录':p.paragraph_format.page_break_before=True
                self.text(p,title);i+=1;continue
            if line.startswith('$$'):
                expression=re.fullmatch(r'\$\$(.+)\$\$',line).group(1)
                match=re.fullmatch(r'(.+?)=\\frac\{([^{}]+)\}\{([^{}]+)\}',expression)
                if not match:raise ValueError('Example renderer supports only named simple fractions')
                p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                math=OxmlElement('m:oMath')
                def run(text):
                    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r
                math.append(run(match[1]+'='));fraction=OxmlElement('m:f')
                for tag,value in [('m:num',match[2]),('m:den',match[3])]:
                    part=OxmlElement(tag);part.append(run(value));fraction.append(part)
                math.append(fraction);p._p.append(math);i+=1;continue
            parts=[]
            while i<len(lines) and lines[i].strip():
                current=lines[i].strip()
                if parts and (current.startswith(('|','#','$$')) or re.fullmatch(r'!\[[^\]]*\]\([^\s)]+\)',current)):break
                parts.append(current);i+=1
            p=self.doc.add_paragraph();self.text(p,'\n'.join(parts))
            if re.fullmatch(r'[图表] \d+','\n'.join(parts)):
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
                for run in p.runs:run.bold=True

def main():
    rt=Runtime(PROJECT);state=rt.read();cp=state['copilot']
    contract=json.loads((PROJECT/'paper/source_contract.json').read_text(encoding='utf-8'))
    plan=json.loads((PROJECT/'paper/chain_plan.json').read_text(encoding='utf-8'))
    registry=json.loads((PROJECT/'paper/reference_registry.json').read_text(encoding='utf-8'))
    paper_source_blocks(PROJECT,cp,plan['sections'],contract)
    doc=make_document();renderer=Renderer(doc,registry)
    for oid in contract['sections']:
        payload=cp['objects'][oid]['payload'];text=(PROJECT/payload['path']).read_text(encoding='utf-8')
        renderer.markdown(section_projection(cp,text,payload.get('source_bindings'),payload.get('structure'),root=PROJECT))
    for item in contract['supplements']:
        if item['kind']=='file_list':
            for relative in item['paths']:
                p=doc.add_paragraph('文件：'+relative);p.paragraph_format.space_after=Pt(2)
                for run in p.runs:run.font.size=Pt(9)
            continue
        p=doc.add_paragraph('源程序：'+item['path']);p.paragraph_format.keep_with_next=True
        content=(PROJECT/item['path']).read_text(encoding='utf-8-sig').replace('\r\n','\n').rstrip('\n')
        doc.add_paragraph(content,style='CUMCMSourceCode')
    target=PROJECT/'paper/历史复现核验报告.docx';doc.save(target)
    report=audit_paper_source(PROJECT,cp,plan['sections'],target,contract)
    save_json(PROJECT/'paper/source-audit-before-render.json',report)
    if not report['passed']:raise RuntimeError(report['errors'])
    metadata={'sha256':sha(target),'milestone':'working','history_class':'working_build','formal_history_eligible':False,
      'artifact_id':'ART-HISTORICAL-PAPER','formal_values_present':True,'run_id':plan['benchmark_ids']['Q2']['run'],
      'source_state_revision':cp['revision'],'scope':'historical_benchmark; not human-final or formal competition paper'}
    save_json(PROJECT/'paper/paper.metadata.json',metadata)
    ai=make_document();ai.core_properties.title='AI 工具使用详情';render=Renderer(ai)
    ai_source=(PROJECT/'paper/AI工具使用详情.md').read_text(encoding='utf-8')
    # The disclosure export contains emphasis/code markers, not executable code.
    # Strip only these inline wrappers so they do not print as literal Markdown.
    ai_source=re.sub(r'\*\*([^\n]+?)\*\*',r'\1',ai_source)
    ai_source=re.sub(r'`([^`\n]+)`',r'\1',ai_source)
    render.markdown(ai_source)
    ai.save(PROJECT/'paper/AI工具使用详情.docx')
    print(json.dumps({'docx':str(target),'source_audit_passed':report['passed'],'source_blocks':report['source_block_count']},ensure_ascii=False))

if __name__=='__main__':main()
