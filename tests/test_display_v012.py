"""v0.1.2 explicit presentation contracts: arithmetic and real Runtime boundaries."""
from __future__ import annotations

import copy
import decimal
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from copilot_display import validate_display_contracts, resolve_metric
from copilot_store import digest


def authority(metrics=None):
    """Synthetic object graph for pure arithmetic tests, not execution evidence."""
    def envelope(ident, key, kind, payload, dependencies):
        return {"id": ident, "key": key, "kind": kind, "status": "verified",
                "payload": payload, "payload_hash": digest(payload),
                "dependencies": dependencies, "files": []}
    run = envelope("run@1", "run", "RunRecord", {"question": "Q1"}, [])
    result = envelope("result@1", "result", "ResultRecord",
                      {"run_id": "run@1", "metrics": metrics if metrics is not None else {"value": 182.8125}},
                      ["run@1"])
    return {"objects": {"run@1": run, "result@1": result},
            "current": {"run": "run@1", "result": "result@1"}}


def contract(raw="182.8125", shown="182.81", **changes):
    result = {"version": "0.1", "result_id": "result@1", "metric_path": "/value",
              "raw_value": raw, "display_value": shown, "format": "decimal",
              "decimal_places": 2, "rounding": "half_even"}
    result.update(changes)
    return result


class DisplayContractTests(unittest.TestCase):
    def validate(self, item=None, cp=None, **kwargs):
        return validate_display_contracts(cp if cp is not None else authority(), ["result@1"],
                                          [item if item is not None else contract()], **kwargs)

    def test_correct_rounding_is_auditable_and_does_not_write_inputs(self):
        cp, contracts = authority(), [contract()]
        original = copy.deepcopy((cp, contracts))
        result = validate_display_contracts(cp, ["result@1"], contracts)
        self.assertEqual((cp, contracts), original)
        binding = result[0]
        self.assertEqual(binding["text"], "182.81")
        self.assertEqual(binding["sources"], [{"result_id": "result@1", "metric_path": "/value", "value": 182.8125}])
        self.assertEqual(binding["derived_value"], "182.81")
        self.assertEqual(binding["rounding_delta"], "-0.0025")
        self.assertEqual(binding["rounding_delta_basis"], "source_metric")
        self.assertEqual(binding["source_payload_hash"], cp["objects"]["result@1"]["payload_hash"])
        contracts[0]["raw_value"] = "999"
        self.assertEqual(binding["display_contract"]["raw_value"], "182.8125")

    def test_false_display_is_rejected_without_relaxing_tolerance(self):
        for shown in ("182.82", "183.00", "999", "182.8100001"):
            with self.subTest(shown=shown), self.assertRaises(ValueError):
                self.validate(contract(shown=shown))

    def test_half_even_and_half_up_are_explicit_and_decimal(self):
        for raw, even, up in ((2.665, "2.66", "2.67"), (2.675, "2.68", "2.68"),
                              (-2.665, "-2.66", "-2.67"), (9.995, "10.00", "10.00")):
            for mode, shown in (("half_even", even), ("half_up", up)):
                with self.subTest(raw=raw, mode=mode):
                    self.assertEqual(self.validate(contract(str(raw), shown, rounding=mode), authority({"value": raw}))[0]["text"], shown)

    def test_zero_and_maximum_supported_decimal_places(self):
        self.assertEqual(self.validate(contract(shown="183", decimal_places=0))[0]["text"], "183")
        self.assertEqual(self.validate(contract(shown="182.812500000000", decimal_places=12))[0]["text"], "182.812500000000")

    def test_percent_has_fixed_fraction_scale_and_delta_in_source_units(self):
        binding = self.validate(contract("0.1828125", "18.28%", format="percent"), authority({"value": 0.1828125}))[0]
        self.assertEqual(binding["rounding_delta"], "-0.0000125")
        self.assertEqual(binding["sources"][0]["value"], 0.1828125)
        self.assertEqual(self.validate(contract("-0.125", "-12.50%", format="percent"), authority({"value": -0.125}))[0]["text"], "-12.50%")

    def test_percent_without_percent_or_with_arbitrary_scale_is_rejected(self):
        for shown in ("18.28", "0.18%", "1828.13%", "18.28 %", "18.28％"):
            with self.subTest(shown=shown), self.assertRaises(ValueError):
                self.validate(contract("0.1828125", shown, format="percent"), authority({"value": 0.1828125}))
        with self.assertRaises(ValueError):
            self.validate(contract(shown="182.81%"))

    def test_raw_value_must_match_named_metric_before_rounding(self):
        for raw in ("182.81", "182.8125001", "999"):
            with self.subTest(raw=raw), self.assertRaises(ValueError):
                self.validate(contract(raw=raw))

    def test_float_serialization_cannot_be_silently_replaced_by_shorter_raw(self):
        value = 0.1 + 0.2
        cp = authority({"value": value})
        with self.assertRaises(ValueError):
            self.validate(contract("0.3", "0.30"), cp)
        self.assertEqual(self.validate(contract(str(value), "0.30"), cp)[0]["sources"][0]["value"], value)

    def test_json_pointer_nested_arrays_and_escaped_keys(self):
        metrics = {"groups": [{"value": 2.665}], "a/b": {"~value": 3.125}, "0": 182.8125}
        cp = authority(metrics)
        for pointer, raw, shown in (("/groups/0/value", "2.665", "2.66"),
                                     ("/a~1b/~0value", "3.125", "3.12"),
                                     ("/0", "182.8125", "182.81")):
            with self.subTest(pointer=pointer):
                self.validate(contract(raw, shown, metric_path=pointer), cp)

    def test_pointer_is_not_an_expression_or_noncanonical_list_index(self):
        metrics = {"groups": [{"value": 2.665}]}
        for pointer in ("groups[0].value", "/groups/01/value", "/groups/-1/value", "/groups/1/value",
                        "/groups/*/value", "/groups/0/missing", "/groups/0/value~2", "", None,
                        "/" + "v" * 1024, "/x" * 33):
            with self.subTest(pointer=pointer), self.assertRaises(ValueError):
                resolve_metric(metrics, pointer)

    def test_bool_text_container_and_nonfinite_metrics_rejected(self):
        for value in (True, False, "182.8125", {"value": 182.8125}, [182.8125], None, float("nan"), float("inf")):
            with self.subTest(value=value), self.assertRaises(ValueError):
                resolve_metric({"value": value}, "/value")

    def test_source_must_be_in_enclosing_claim_dependencies(self):
        with self.assertRaises(ValueError):
            validate_display_contracts(authority(), [], [contract()])
        with self.assertRaises(ValueError):
            self.validate(contract(result_id="result@999"))

    def test_source_must_exist_and_be_verified_result(self):
        for alteration in ("missing", "kind", "generated", "stale", "superseded"):
            cp = authority()
            if alteration == "missing":
                del cp["objects"]["result@1"]
            elif alteration == "kind":
                cp["objects"]["result@1"]["kind"] = "ArtifactRecord"
            else:
                cp["objects"]["result@1"]["status"] = alteration
            with self.subTest(alteration=alteration), self.assertRaises(ValueError):
                self.validate(cp=cp)

    def test_source_current_id_and_payload_hash_must_match(self):
        for alteration in ("current", "id", "payload_hash", "metric"):
            cp = authority()
            if alteration == "current":
                cp["current"]["result"] = "result@2"
            elif alteration == "id":
                cp["objects"]["result@1"]["id"] = "wrong"
            elif alteration == "payload_hash":
                cp["objects"]["result@1"]["payload_hash"] = "f" * 64
            else:
                cp["objects"]["result@1"]["payload"]["metrics"]["value"] = 999
            with self.subTest(alteration=alteration), self.assertRaises(ValueError):
                self.validate(cp=cp)

    def test_ancestor_stale_noncurrent_missing_hash_and_cycle_are_rejected(self):
        for alteration in ("stale", "current", "missing", "payload_hash", "cycle"):
            cp = authority()
            if alteration == "stale":
                cp["objects"]["run@1"]["status"] = "stale"
            elif alteration == "current":
                cp["current"]["run"] = "run@2"
            elif alteration == "missing":
                del cp["objects"]["run@1"]
            elif alteration == "payload_hash":
                cp["objects"]["run@1"]["payload_hash"] = "forged"
            else:
                cp["objects"]["run@1"]["dependencies"] = ["result@1"]
            with self.subTest(alteration=alteration), self.assertRaises(ValueError):
                self.validate(cp=cp)

    def test_unknown_fields_cannot_be_used_as_tolerance_or_exemption(self):
        for name, value in (("tolerance", 1000), ("relative_tolerance", 1), ("scale", 1e6),
                            ("offset", 816.1875), ("number_kind", "year"), ("constant", True),
                            ("source_verified", True), ("format_expression", "round(x, 2)")):
            with self.subTest(name=name), self.assertRaises(ValueError):
                self.validate(contract(**{name: value}))

    def test_all_fields_are_required(self):
        for field in contract():
            item = contract()
            del item[field]
            with self.subTest(field=field), self.assertRaises(ValueError):
                self.validate(item)

    def test_unknown_enum_and_noninteger_or_unbounded_precision_are_rejected(self):
        mutations = [{"format": value} for value in ("scientific", "ratio", [], 0)]
        mutations += [{"rounding": value} for value in ("up", "down", "any", [], 0)]
        mutations += [{"version": value} for value in ("0.2", 0.1, None)]
        mutations += [{"decimal_places": value} for value in (True, False, -1, 13, "2", 2.0, None)]
        for change in mutations:
            with self.subTest(change=change), self.assertRaises(ValueError):
                self.validate(contract(**change))

    def test_raw_decimal_strings_are_finite_bounded_and_ascii(self):
        for raw in ("NaN", "Infinity", "1e1001", "1e-1001", "1,82.8125", "１８２.８１２５", " 182.8125 ",
                    "1" * 513, 182.8125, True, None):
            with self.subTest(raw=raw), self.assertRaises(ValueError):
                self.validate(contract(raw=raw))

    def test_nonjson_contract_parameters_raise_contract_error(self):
        for field, value in (("raw_value", decimal.Decimal("182.8125")),
                             ("display_value", {"182.81"}), ("metric_path", object())):
            with self.subTest(field=field), self.assertRaises(ValueError):
                self.validate(contract(**{field: value}))

    def test_display_string_must_match_declared_format_exactly(self):
        for shown in ("182.810", "+182.81", "1.8281e2", " 182.81", "１８２.８１", "182.81 m", 182.81):
            with self.subTest(shown=shown), self.assertRaises(ValueError):
                self.validate(contract(shown=shown))

    def test_nonzero_to_zero_is_rejected_but_real_zero_is_valid(self):
        with self.assertRaises(ValueError):
            self.validate(contract("0.0001", "0.00"), authority({"value": 0.0001}))
        self.assertEqual(self.validate(contract("0", "0.00"), authority({"value": 0}))[0]["text"], "0.00")

    def test_rounding_is_independent_of_ambient_decimal_context(self):
        with decimal.localcontext() as context:
            context.prec = 2
            context.rounding = decimal.ROUND_DOWN
            context.traps[decimal.Inexact] = True
            self.assertEqual(self.validate()[0]["text"], "182.81")

    def test_source_error_callback_can_reject_physical_drift(self):
        with self.assertRaisesRegex(ValueError, "hash changed"):
            self.validate(source_errors=lambda ident: ["hash changed for " + ident])

    def test_callback_is_trusted_error_list_not_a_flag(self):
        for callback in (True, {"verified": True}, lambda ident: True, lambda ident: None,
                         lambda ident: [False], lambda ident: ("error",)):
            with self.subTest(callback=callback), self.assertRaises(ValueError):
                self.validate(source_errors=callback)

    def test_callback_checks_each_source_once_without_mutation(self):
        seen = []
        def check(ident):
            seen.append(ident)
            return []
        bindings = validate_display_contracts(authority(), ["result@1"],
            [contract(), contract(shown="182.812", decimal_places=3)], source_errors=check)
        self.assertEqual(len(bindings), 2)
        self.assertEqual(seen, ["result@1"])

    def test_duplicate_or_excessive_contracts_are_rejected(self):
        for contracts in ([contract(), contract()], [contract()] * 257, {}, None):
            with self.subTest(count=len(contracts) if contracts is not None else None), self.assertRaises(ValueError):
                validate_display_contracts(authority(), ["result@1"], contracts)

    def test_empty_contracts_preserve_exact_number_path(self):
        self.assertEqual(validate_display_contracts(authority(), ["result@1"], []), [])

    def test_malformed_authority_is_a_clear_validation_error(self):
        for cp in (None, [], {"objects": None}, {"objects": []}, {"objects": {}, "current": []}):
            with self.subTest(cp=cp), self.assertRaises(ValueError):
                self.validate(cp=cp if cp is not None else [])

    def test_year_and_formula_labels_are_not_exemptions(self):
        for item in (contract("2026", "2026", decimal_places=0),
                     contract(shown="999", decimal_places=0, kind="formula_constant"),
                     contract(shown="182.81", year=2026)):
            with self.subTest(item=item), self.assertRaises(ValueError):
                self.validate(item)


class DisplayRuntimeBoundaryTests(unittest.TestCase):
    def setUp(self):
        # Reuse the frozen regression fixture without modifying it: actual code
        # reads the runtime parameter context, and a separate checker adds n+n.
        from test_copilot_runtime import make_project
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.rt, self.ids, self.register = make_project(self.root, count=1)
        parameters = copy.deepcopy(self.rt.read()["copilot"]["objects"][self.ids["params"]]["payload"])
        parameters["entries"][0]["current_value"] = 91.40625
        self.ids["params"] = self.register("ParameterSet", "params.Q1", parameters, [self.ids["model"]])
        self.run_id = self.rt.execute(self.revision(), "Q1", ["{python}", "solver.py"], ["result.json"],
            dependencies=[self.ids[k] for k in ("model", "params", "data", "code", "plan")])["result"]["object_id"]
        checked = self.rt.validate_run(self.revision(), self.run_id, "checker.py",
            ["{python}", "{checker}", "{run}", "{report}"])["result"]
        self.assertTrue(checked["passed"], checked)
        self.result_id = checked["result_id"]
        self.item = contract(result_id=self.result_id)

    def revision(self):
        return self.rt.read()["copilot"]["revision"]

    def validate_current(self):
        from copilot_runtime import object_errors
        cp = self.rt.read()["copilot"]
        return validate_display_contracts(cp, [self.result_id], [self.item],
            source_errors=lambda ident: object_errors(self.root, cp, ident))

    def test_real_checked_metric_supports_rounding_without_authority_write(self):
        before = self.rt.store.path.read_bytes()
        original = copy.deepcopy(self.rt.read()["copilot"]["objects"][self.result_id])
        self.assertEqual(original["payload"]["metrics"]["value"], 182.8125)
        self.assertEqual(self.validate_current()[0]["text"], "182.81")
        self.assertEqual(self.rt.store.path.read_bytes(), before)
        self.assertEqual(self.rt.read()["copilot"]["objects"][self.result_id], original)

    def test_actual_output_file_drift_rejected_without_reconcile_write(self):
        cp = self.rt.read()["copilot"]
        output = cp["objects"][self.run_id]["payload"]["outputs"][0]["path"]
        before = self.rt.store.path.read_bytes()
        (self.root / output).write_text('{"value":999}', encoding="utf-8")
        self.assertEqual(self.rt.read()["copilot"]["objects"][self.result_id]["status"], "verified")
        with self.assertRaises(ValueError):
            self.validate_current()
        self.assertEqual(self.rt.store.path.read_bytes(), before)

    def test_parameter_revision_invalidates_display_even_when_old_arithmetic_matches(self):
        parameters = copy.deepcopy(self.rt.read()["copilot"]["objects"][self.ids["params"]]["payload"])
        parameters["entries"][0]["current_value"] = 92
        self.register("ParameterSet", "params.Q1", parameters, [self.ids["model"]])
        cp = self.rt.read()["copilot"]
        self.assertEqual(cp["objects"][self.result_id]["status"], "stale")
        with self.assertRaises(ValueError):
            validate_display_contracts(cp, [self.result_id], [self.item])


class DisplaySectionIntegrationRedTeam(DisplayRuntimeBoundaryTests):
    # Inherit the real execution setup, but keep the three original boundary
    # tests on their original class so this independent audit does not recount
    # them as additional observations.
    test_real_checked_metric_supports_rounding_without_authority_write = None
    test_actual_output_file_drift_rejected_without_reconcile_write = None
    test_parameter_revision_invalidates_display_even_when_old_arithmetic_matches = None

    def checked_claim(self):
        cp = self.rt.read()["copilot"]
        run = cp["objects"][self.run_id]["payload"]
        self.claim_text = "The computed length is 182.81 m."
        payload = {"claim_id": "display.audit", "claim_type": "numerical", "claim": self.claim_text,
                   "paper_anchor": "results", "formal_run_id": self.run_id,
                   "requirement_ids": ["REQ-Q1-001"], "data_sources": [run["data_hash"]],
                   "code_locations": ["solver.py"], "tables": [run["outputs"][0]["path"]],
                   "validation_evidence": ["Q1.run.formal"], "limitations": ["Synthetic deterministic test only."],
                   "display_contracts": [self.item]}
        self.claim_id = self.rt.claim(self.revision(), payload, [self.result_id])["result"]["claim_id"]
        return self.claim_text + " [[claim:" + self.claim_id + "]]"

    def test_figure_reference_cannot_consume_whitespace_exponent_prefix(self):
        from copilot_delivery import Delivery
        from PIL import Image
        marked = self.checked_claim()
        Image.new("RGB", (8, 8), "navy").save(self.root / "figure.png")
        artifact = self.register("ArtifactRecord", "figure.audit", {"path": "figure.png", "artifact_type": "figure"}, [self.result_id], ["figure.png"])
        text = "[[figure:1]]\n\n图1 e3 m 的计算误差。\n\n" + marked
        (self.root / "section.md").write_text(text, encoding="utf-8")
        with self.assertRaises(ValueError):
            Delivery(self.root).section(self.revision(), "paper.audit", "section.md", [self.claim_id],
                structure=[{"kind": "figure", "number": 1, "artifact_id": artifact}])

    def table_source(self):
        from copilot_delivery import Delivery
        from docx import Document
        marked = self.checked_claim()
        (self.root / "table-section.md").write_text("| Finding |\n| --- |\n| " + marked + " |", encoding="utf-8")
        section_id = Delivery(self.root).section(self.revision(), "paper.table.audit", "table-section.md", [self.claim_id])["result"]["section_id"]
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "Finding"
        table.cell(1, 0).text = self.claim_text
        return doc, table, section_id

    def table_audit(self, doc, section_id):
        from copilot_paper_source import audit_paper_source
        doc.save(self.root / "actual.docx")
        return audit_paper_source(self.root, self.rt.read()["copilot"], [section_id], "actual.docx",
                                 {"version": "0.1", "sections": [section_id]})

    def test_hidden_table_text_cannot_establish_visible_claim_coverage(self):
        doc, table, section_id = self.table_source()
        self.assertTrue(self.table_audit(doc, section_id)["passed"])
        table.cell(1, 0).paragraphs[0].runs[0].font.hidden = True
        result = self.table_audit(doc, section_id)
        self.assertFalse(result["passed"], result)

    def test_visible_content_control_in_table_cannot_disappear_from_readback(self):
        from docx.oxml import OxmlElement
        doc, table, section_id = self.table_source()
        self.assertTrue(self.table_audit(doc, section_id)["passed"])
        control, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
        paragraph, run, text = OxmlElement("w:p"), OxmlElement("w:r"), OxmlElement("w:t")
        text.text = "A fabricated additional result is 999 m."
        run.append(text); paragraph.append(run); content.append(paragraph); control.append(content)
        table.cell(1, 0)._tc.append(control)
        result = self.table_audit(doc, section_id)
        self.assertFalse(result["passed"], result)

    def test_docx_whitespace_cannot_split_a_bound_numeric_literal(self):
        doc, table, section_id = self.table_source()
        self.assertTrue(self.table_audit(doc, section_id)["passed"])
        table.cell(1, 0).text = self.claim_text.replace("182.81", "18 2.81")
        result = self.table_audit(doc, section_id)
        self.assertFalse(result["passed"], result)


class DisplayComparisonTests(unittest.TestCase):
    def make_case(self, statement=None):
        from test_copilot_runtime import make_project, CHECKER
        import copilot_domain as domain
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.rt, self.ids, register = make_project(self.root, count=1)
        parameters = copy.deepcopy(self.rt.read()["copilot"]["objects"][self.ids["params"]]["payload"])
        parameters["entries"][0]["current_value"] = 0.0245
        self.ids["params"] = register("ParameterSet", "params.Q1", parameters, [self.ids["model"]])
        if statement is not None:
            # The independent checker actually evaluates the raw threshold.
            # The optional extra-number case is an intentionally adversarial
            # checker statement, never a valid numerical evidence exemption.
            checker = CHECKER.replace("ok=actual==sum([n,n])", "ok=actual==sum([n,n])\nstatement=" + repr(statement) + " if actual < 0.05 else 'Threshold condition is false.'")
            checker = checker.replace("'metrics':{'value':actual}", "'metrics':{'value':actual,'statement':statement}")
            (self.root / "checker.py").write_text(checker, encoding="utf-8")
            plan = copy.deepcopy(self.rt.read()["copilot"]["objects"][self.ids["plan"]]["payload"])
            plan["checker"]["sha256"] = domain.sha256_file(self.root / "checker.py")
            self.ids["plan"] = register("ValidationPlan", "plan.Q1", plan, [self.ids["model"]])
        self.run_id = self.rt.execute(self.revision(), "Q1", ["{python}", "solver.py"], ["result.json"],
            dependencies=[self.ids[k] for k in ("model", "params", "data", "code", "plan")])["result"]["object_id"]
        checked = self.rt.validate_run(self.revision(), self.run_id, "checker.py",
            ["{python}", "{checker}", "{run}", "{report}"])["result"]
        self.assertTrue(checked["passed"], checked)
        self.result_id = checked["result_id"]
        self.assertEqual(self.rt.read()["copilot"]["objects"][self.result_id]["payload"]["metrics"]["value"], 0.049)
        self.serial = 0

    def revision(self):
        return self.rt.read()["copilot"]["revision"]

    def claim(self, text, kind="numerical"):
        self.serial += 1
        run = self.rt.read()["copilot"]["objects"][self.run_id]["payload"]
        payload = {"claim_id": "display.comparison." + str(self.serial), "claim_type": kind, "claim": text,
                   "paper_anchor": "results", "formal_run_id": self.run_id,
                   "requirement_ids": ["REQ-Q1-001"], "data_sources": [run["data_hash"]],
                   "code_locations": ["solver.py"], "tables": [run["outputs"][0]["path"]],
                   "validation_evidence": ["Q1.run.formal"], "limitations": ["Synthetic deterministic threshold fixture only."],
                   "display_contracts": [contract("0.049", "0.05", result_id=self.result_id)]}
        return self.rt.claim(self.revision(), payload, [self.result_id])["result"]

    def test_approximate_display_preserves_normal_noncomparison_claim(self):
        self.make_case()
        result = self.claim("The computed value is approximately 0.05.")
        self.assertEqual(result["status"], "verified")
        self.assertEqual(result["assurance"]["comparison_status"], "not_present")

    def test_rounded_inequalities_require_checker_conclusion(self):
        self.make_case()
        for statement in ("The computed value is >= 0.05.", "The computed value is ＞＝ 0.05.",
                          "计算值至少为 0.05。", "The computed value meets threshold 0.05."):
            with self.subTest(statement=statement), self.assertRaises(ValueError):
                self.claim(statement)

    def test_relabelled_comparison_is_not_verified(self):
        from copilot_runtime import usable
        self.make_case()
        for kind in ("model", "method", "theory", "limitation"):
            with self.subTest(kind=kind):
                result = self.claim("The computed value is >= 0.05.", kind)
                self.assertEqual(result["status"], "generated")
                self.assertFalse(result["assurance"]["verified"])
                self.assertFalse(usable(self.root, self.rt.read()["copilot"], result["claim_id"], verified=True))

    def test_true_raw_threshold_computed_by_independent_checker_can_bind_conclusion(self):
        text = "The computed value is below threshold 0.05."
        self.make_case(text)
        result = self.claim(text)
        self.assertEqual(result["status"], "verified")
        self.assertEqual(result["assurance"]["comparison_status"], "bound_to_checker_statement")
        self.assertEqual(result["assurance"]["statement_bindings"][0]["value"], text)

    def test_checker_statement_cannot_exempt_an_additional_unbound_number(self):
        text = "The computed value is below threshold 0.05 with fabricated gain 999."
        self.make_case(text)
        with self.assertRaises(ValueError):
            self.claim(text)

    def test_another_result_statement_cannot_certify_rounded_source_comparison(self):
        # Pure inference fixture: all envelopes are current, but the statement
        # describes a different Result. No real execution is claimed here.
        from copilot_runtime import claim_assurance
        text = "The computed value is >= 0.05."
        cp = authority({"value": 0.049})
        other = copy.deepcopy(cp["objects"]["result@1"])
        other.update(id="other@1", key="other")
        other["payload"]["metrics"] = {"value": 0.051, "statement": text}
        other["payload_hash"] = digest(other["payload"])
        cp["objects"]["other@1"] = other
        cp["current"]["other"] = "other@1"
        result = claim_assurance(cp, {"claim": text, "display_contracts": [contract("0.049", "0.05")]},
                                 ["result@1", "other@1"])
        self.assertFalse(result["verified"], result)


if __name__ == "__main__":
    unittest.main(verbosity=2)
