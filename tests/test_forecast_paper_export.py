"""Actual known-case runs, native source export and dependency invalidation."""
import copy
import importlib.util
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from copilot_paper_export import Renderer, export, native_math
from copilot_runtime import Runtime
from copilot_view import snapshot

def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
    return module

class ForecastPaperTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.fixture = Path(cls.temp.name) / "fixture"
        cls.report = load("forecast_paper_example", ROOT / "examples/forecast_paper/run_example.py").run_example(cls.fixture)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name) / "project"
        shutil.copytree(self.fixture, self.root)
        self.rt = Runtime(self.root); self.ids = self.report["objects"]

    def revision(self): return self.rt.read()["copilot"]["revision"]

    def test_actual_metrics_and_receipts_with_read_only_comparison(self):
        before = self.rt.store.path.read_bytes()
        view = snapshot(self.root)
        group = view["comparisons"][0]
        self.assertTrue(group["comparable"])
        self.assertEqual([(r["value"], r["unit"], r["checked"]) for r in group["rows"]], [(0, "1", True), (5, "1", True)])
        cp = self.rt.read()["copilot"]
        self.assertEqual(cp["objects"][self.ids["result"]]["status"], "verified")
        self.assertTrue(any(f["path"].endswith("sealed-report.json") for f in cp["objects"][self.ids["validation"]]["files"]))
        self.assertTrue(list((self.root / ".copilot/runs").rglob("receipt.json")))
        self.assertEqual(before, self.rt.store.path.read_bytes())

    def test_native_chinese_docx_table_formula_figure_and_audit(self):
        result = self.report["export"]
        with zipfile.ZipFile(self.root / result["docx"]) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
            self.assertIn("合成预测与论文导出验证", xml)
            self.assertIn("<w:tbl>", xml); self.assertIn("<m:oMath>", xml); self.assertIn("<m:f>", xml)
            self.assertIn("<w:drawing>", xml); self.assertNotIn("\\frac", xml)
        audit = json.loads((self.root / result["source_audit"]).read_text(encoding="utf-8"))
        self.assertTrue(audit["passed"])
        self.assertFalse(result["submission_ready"])
        self.assertEqual(result["visual_review"], "not_performed")

    def test_long_numeric_table_preserves_digits_and_allocates_readable_width(self):
        from docx import Document
        from docx.shared import Pt
        rows = [["起点或汇总", "步长", "真值（kW）", "回归（kW）", "基线（kW）", "回归绝对误差（kW）", "基线绝对误差（kW）"],
                ["十六时", "一小时", "135.0", "131.538461538462", "129.0", "3.461538461538", "6.0"],
                ["整体 MAE", "—", "—", "—", "—", "0.776923076923", "5.0"]]
        renderer = Renderer(self.root); renderer.table(rows)
        target = self.root / "wide-numbers.docx"; renderer.doc.save(target)
        doc = Document(target); table = doc.tables[0]
        self.assertEqual([[cell.text for cell in row.cells] for row in table.rows], rows)
        available = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
        self.assertLessEqual(sum(column.width for column in table.columns), available + Pt(1))
        self.assertGreaterEqual(table.columns[3].width, Pt(110))
        self.assertLess(table.columns[0].width, table.columns[3].width)
        self.assertGreaterEqual(table.cell(1, 3).paragraphs[0].runs[0].font.size, Pt(9))

    def test_unreadably_wide_numeric_table_is_refused_without_rounding(self):
        renderer = Renderer(self.root)
        with self.assertRaisesRegex(ValueError, "不会自动删减数值精度"):
            renderer.table([["值"] * 12, ["12345678901234567890.123456789"] * 12])

    def test_documented_public_cli_exports_native_file(self):
        completed = subprocess.run([sys.executable, "-X", "utf8", str(ROOT/"scripts/copilot.py"),
            "--workspace", str(self.root), "paper-export", "--expected-revision", str(self.revision()),
            "--contract", "paper/source_contract.json", "--output", "paper/命令行导出.docx"],
            capture_output=True, text=True, encoding="utf-8", timeout=60)
        self.assertEqual(completed.returncode, 0, completed.stdout+completed.stderr)
        self.assertTrue(json.loads(completed.stdout)["ok"])
        self.assertTrue((self.root/"paper/命令行导出.docx").is_file())

    def test_output_create_only_and_revision_are_enforced_before_write(self):
        output = self.root / self.report["export"]["docx"]
        saved = output.read_bytes(); state = self.rt.store.path.read_bytes()
        with self.assertRaises(ValueError):
            export(self.root, "paper/source_contract.json", self.report["export"]["docx"], expected_revision=self.revision())
        with self.assertRaises(ValueError):
            export(self.root, "paper/source_contract.json", "paper/new.docx", expected_revision=self.revision()-1)
        self.assertEqual(saved, output.read_bytes()); self.assertEqual(state, self.rt.store.path.read_bytes())
        self.assertFalse((self.root / "paper/new.docx").exists())

    def test_parameter_change_invalidates_result_section_and_docx(self):
        payload = copy.deepcopy(self.rt.read()["copilot"]["objects"][self.ids["params"]]["payload"])
        payload["version"] = 2; payload["entries"][0]["current_value"] = 5
        self.rt.register(self.revision(), "ParameterSet", "params.Q1", payload, dependencies=[self.ids["model"]])
        view = snapshot(self.root)
        for oid in (self.ids["result"], self.ids["section"], self.report["export"]["artifact_id"]):
            self.assertEqual(view["objects"][oid]["effective_status"], "stale")
        self.assertFalse(view["comparisons"][0]["comparable"])
        with self.assertRaises(ValueError):
            export(self.root, "paper/source_contract.json", "paper/after-change.docx", expected_revision=self.revision())
        self.assertFalse((self.root / "paper/after-change.docx").exists())

    def test_same_revision_source_drift_blocks_export(self):
        (self.root / "paper/results.md").write_text("用户修改了原稿", encoding="utf-8")
        state = self.rt.store.path.read_bytes()
        with self.assertRaises(ValueError):
            export(self.root, "paper/source_contract.json", "paper/drift.docx", expected_revision=self.revision())
        self.assertFalse((self.root / "paper/drift.docx").exists())
        self.assertEqual(state, self.rt.store.path.read_bytes())

    def test_wrong_comparison_hash_or_fields_cannot_register(self):
        cp = self.rt.read()["copilot"]; before = self.rt.store.path.read_bytes()
        for error in ("hash", "unknown"):
            payload = copy.deepcopy(cp["objects"][self.ids["model"]]["payload"])
            comparison = payload["data_contract"]["comparison"]
            if error == "hash": comparison["source_files"][0]["sha256"] = "0"*64
            else: comparison["variants"][0]["context"]["unknown"] = "ignored?"
            with self.assertRaises(ValueError):
                self.rt.register(self.revision(), "ModelSpec", "model.Q1", payload, dependencies=[self.ids["contract"]])
            self.assertEqual(before, self.rt.store.path.read_bytes())

    def test_independent_checker_rejects_forged_prediction_and_scaler(self):
        cp = self.rt.read()["copilot"]
        run = cp["objects"][self.ids["run"]]["payload"]
        directory = self.root / Path(run["outputs"][0]["path"]).parent
        checker = load("forecast_independent_checker", ROOT / "examples/forecast_paper/checker.py")
        result = directory / "result.json"; original = json.loads(result.read_text(encoding="utf-8"))
        for field in ("prediction", "scaler"):
            modified = copy.deepcopy(original)
            if field == "prediction": modified["forecasts"][0]["prediction"] += 9
            else: modified["forecasts"][0]["preprocessing"]["mean"][0] += 9
            result.write_text(json.dumps(modified), encoding="utf-8")
            self.assertTrue(any(row["status"] == "fail" for row in checker.check(directory)["checks"]))

    def test_pdf_backend_failure_preserves_audited_word_without_certification(self):
        with patch("copilot_paper_export.convert_pdf", side_effect=ValueError("unavailable test backend")):
            with self.assertRaisesRegex(ValueError, "unavailable"):
                export(self.root, "paper/source_contract.json", "paper/backend-failure.docx", expected_revision=self.revision(), pdf="paper/backend-failure.pdf")
        self.assertTrue((self.root / "paper/backend-failure.docx").is_file())
        self.assertFalse((self.root / "paper/backend-failure.pdf").exists())
        artifacts = [o for o in self.rt.read()["copilot"]["objects"].values() if o["payload"].get("path") == "paper/backend-failure.docx"]
        self.assertEqual(len(artifacts), 1); self.assertEqual(artifacts[0]["status"], "generated")
        self.assertEqual(artifacts[0]["payload"]["visual_review"], "not_performed")

    def test_complex_math_refused_and_invalid_pdf_path_leaves_no_output(self):
        with self.assertRaises(ValueError): native_math(r"y=\sum_{i=1}^{N}x_i")
        with self.assertRaises(ValueError):
            export(self.root, "paper/source_contract.json", "paper/invalid.docx", expected_revision=self.revision(), pdf="../../escape.pdf")
        self.assertFalse((self.root / "paper/invalid.docx").exists())
