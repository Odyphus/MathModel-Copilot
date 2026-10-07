"""Independent v0.1.2 acceptance/adversarial tests, with actual checked runs.

Synthetic source/rule/review fixtures do not assert an external human signoff.
MATHMODEL_TEST_REPO allows the unchanged tests to run against a read-only base.
"""
from __future__ import annotations

import copy
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(os.environ.get("MATHMODEL_TEST_REPO", Path(__file__).resolve().parents[1])).resolve()
sys.path[:0] = [str(ROOT / "scripts"), str(ROOT / "tests")]

from test_copilot_runtime import make_project
import test_copilot_delivery as delivery_fixture
import copilot_domain as domain
from copilot_delivery import Delivery, submission_status, audit_delivery_zip
from copilot_runtime import usable, object_errors, _put


class V012RedTeamTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.rt, self.ids, self.reg = make_project(self.root, count=1)
        self.delivery = Delivery(self.root)
        self.serial = 0

    def rev(self):
        return self.rt.read()["copilot"]["revision"]

    def write(self, path, value):
        target = self.root / path
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(value if isinstance(value, str) else json.dumps(value), encoding="utf-8")

    def params(self, value):
        payload = copy.deepcopy(self.rt.read()["copilot"]["objects"][self.ids["params"]]["payload"])
        payload["version"] += 1
        payload["entries"][0]["current_value"] = value
        self.ids["params"] = self.reg("ParameterSet", "params.Q1", payload, [self.ids["model"]])

    def checked(self, value=10):
        if value != 10:
            self.params(value / 2)
        run = self.rt.execute(self.rev(), "Q1", ["{python}", "solver.py"], ["result.json"],
            dependencies=[self.ids[key] for key in ("model", "params", "data", "code", "plan")])["result"]["object_id"]
        checked = self.rt.validate_run(self.rev(), run, "checker.py",
            ["{python}", "{checker}", "{run}", "{report}"])["result"]
        self.assertTrue(checked["passed"], checked)
        result = checked["result_id"]
        actual = self.rt.read()["copilot"]["objects"][result]["payload"]["metrics"]["value"]
        self.assertEqual(actual, value)
        return run, result

    def payload(self, run, text, contracts=None, kind="numerical"):
        self.serial += 1
        obj = self.rt.read()["copilot"]["objects"][run]["payload"]
        payload = {"claim_id": "C" + str(self.serial), "claim": text, "claim_type": kind,
            "paper_anchor": "Q1 results", "formal_run_id": run,
            "requirement_ids": ["REQ-Q1-001"], "data_sources": [obj["data_hash"]],
            "code_locations": ["solver.py"], "tables": [obj["outputs"][0]["path"]],
            "validation_evidence": ["Q1.run.formal"],
            "limitations": ["Synthetic deterministic fixture only."]}
        if contracts is not None:
            payload["display_contracts"] = contracts
        return payload

    def claim(self, run, result, text="Computed length is 10 m.", contracts=None, kind="numerical"):
        response = self.rt.claim(self.rev(), self.payload(run, text, contracts, kind), [result])["result"]
        cid = response["claim_id"]
        cp = self.rt.read()["copilot"]
        self.assertTrue(usable(self.root, cp, cid, verified=True), response)
        return cid

    def contract(self, result, raw="182.8125", shown="182.81", **updates):
        value = {"version": "0.1", "result_id": result, "metric_path": "/value",
                 "raw_value": raw, "display_value": shown, "format": "decimal",
                 "decimal_places": 2, "rounding": "half_even"}
        value.update(updates)
        return value

    def section(self, text, claims, **kwargs):
        self.serial += 1
        path = "sections/section-" + str(self.serial) + ".md"
        self.write(path, text)
        result = self.delivery.section(self.rev(), "paper." + str(self.serial), path, claims, **kwargs)["result"]
        sid = result["section_id"]
        self.assertTrue(usable(self.root, self.rt.read()["copilot"], sid, verified=True), result)
        return sid

    @staticmethod
    def marked(claim, text="Computed length is 10 m."):
        return text + " [[claim:" + claim + "]]"

    def artifact(self, kind="figure"):
        path = "artifacts/" + kind + ".txt"
        self.write(path, "Synthetic " + kind + " artifact; no inferred numerical finding.")
        return self.reg("ArtifactRecord", "artifact." + kind,
                        {"path": path, "artifact_type": kind})

    def refresh_synthetic_paper_bindings(self):
        """Only test-fixture bookkeeping; never a real render or human signoff."""
        sha = domain.sha256_file(self.root / "paper.docx")
        for path in ("paper.metadata.json", "visual.json", "anonymity.json", "content.json"):
            data = json.loads((self.root / path).read_text(encoding="utf-8"))
            if path == "paper.metadata.json":
                data["sha256"] = sha.lower()
            elif path == "visual.json":
                data["source_docx_sha256"] = sha.lower()
            else:
                data["files"][0]["sha256"] = sha
            self.write(path, data)

    def test_plain_and_outline_headings_preserve_legitimate_result(self):
        run, result = self.checked()
        claim = self.claim(run, result)
        for heading in ("# 模型结果", "# 1 模型结果", "## 2.1 模型结果", "### 3.2.1 模型结果", "# 1. 模型结果"):
            with self.subTest(heading=heading):
                self.section(heading + "\n\n" + self.marked(claim), [claim])

    def test_heading_result_numbers_are_not_outline_exemptions(self):
        run, result = self.checked()
        claim = self.claim(run, result)
        for heading in ("# 结果为 999 m", "# 1 结果为 999 m", "# 99 m", "# 99 m 的计算误差",
                        "# 99 m/s 的运行速度", "# 99 kg·m 的力矩", "# 99 次计算", "# 99 台设备", "# 99% 的准确率",
                        "# 99 ｍ 的计算误差", "# 99 ㎏ 的载荷"):
            with self.subTest(heading=heading):
                with self.assertRaises(ValueError):
                    self.section(heading + "\n\n" + self.marked(claim), [claim])

    def test_heading_or_outline_does_not_exempt_body_value(self):
        run, result = self.checked()
        claim = self.claim(run, result)
        for text in ("# 1 模型结果\n\nFalse value is 999 m. " + self.marked(claim),
                     "# 1 模型结果\n\n1. False value is 999 m. " + self.marked(claim)):
            with self.subTest(text=text):
                with self.assertRaises(ValueError):
                    self.section(text, [claim])

    def test_declared_figure_table_formula_references_are_usable(self):
        run, result = self.checked()
        claim = self.claim(run, result)
        structure = [{"kind": "figure", "number": 1, "artifact_id": self.artifact("figure")},
                     {"kind": "table", "number": 2, "artifact_id": self.artifact("table")},
                     {"kind": "formula", "number": 3, "source_id": self.ids["model"], "source_path": "/formulas/0/expression"}]
        text = "[[figure:1]]\n\n[[table:2]]\n\n[[formula:3]]\n\n图 1、表2与式（3）给出模型。\n\n" + self.marked(claim)
        sid = self.section(text, [claim], structure=structure)
        self.assertTrue(set(item.get("artifact_id", item.get("source_id")) for item in structure)
                        <= set(self.rt.read()["copilot"]["objects"][sid]["dependencies"]))

    def test_undefined_and_falsely_typed_structure_cannot_hide_values(self):
        run, result = self.checked()
        claim = self.claim(run, result)
        artifact = self.artifact("figure")
        cases = [("图999显示结果。", []),
                 ("[[figure:1]]\n\n图1结果为999 m。", [{"kind":"figure", "number":1, "artifact_id":artifact}]),
                 ("[[figure:1]]\n\n图1e3 m 的计算误差。", [{"kind":"figure", "number":1, "artifact_id":artifact}]),
                 ("[[figure:1]]\n\n图1e-3 m 的计算误差。", [{"kind":"figure", "number":1, "artifact_id":artifact}]),
                 ("[[figure:1]]\n\n图1 e3 m 的计算误差。", [{"kind":"figure", "number":1, "artifact_id":artifact}]),
                 ("[[figure:1]]\n\n图1 e 3 m 的计算误差。", [{"kind":"figure", "number":1, "artifact_id":artifact}]),
                 ("[[figure:1]]\n\n图1𝑒3 m 的计算误差。", [{"kind":"figure", "number":1, "artifact_id":artifact}]),
                 ("[[table:1]]", [{"kind":"table", "number":1, "artifact_id":artifact}]),
                 ("[[figure:1]]", [{"kind":"figure", "number":1, "artifact_id":self.ids["model"]}]),
                 ("[[figure:1]]", [{"kind":"figure", "number":True, "artifact_id":artifact}])]
        for text, structure in cases:
            with self.subTest(text=text, structure=structure):
                with self.assertRaises(ValueError):
                    self.section(text + "\n\n" + self.marked(claim), [claim], structure=structure)

    def test_source_notes_are_derived_from_locked_sources(self):
        from copilot_section import source_block_text, section_projection
        run, result = self.checked()
        claim = self.claim(run, result)
        delivery_fixture.DeliveryTests.create_pack(self)
        cp = self.rt.read()["copilot"]
        bindings = [{"id":"year", "kind":"year", "source_id":cp["current"]["rules.lock"], "source_path":"/problem_year"},
                    {"id":"parameter", "kind":"parameter", "source_id":self.ids["params"], "source_path":"/entries/0/current_value"},
                    {"id":"formula", "kind":"formula", "source_id":self.ids["model"], "source_path":"/formulas/0/expression"}]
        text = "\n\n".join("[[source:" + item["id"] + "]]" for item in bindings) + "\n\n" + self.marked(claim)
        sid = self.section(text, [claim], source_bindings=bindings)
        projection = section_projection(cp, text, bindings)
        for item in bindings:
            self.assertIn(source_block_text(cp, item), projection)
        self.assertNotIn("[[source:", projection)
        self.assertTrue(set(item["source_id"] for item in bindings) <= set(self.rt.read()["copilot"]["objects"][sid]["dependencies"]))

    def test_source_label_cannot_exempt_arbitrary_text_or_wrong_paths(self):
        run, result = self.checked()
        claim = self.claim(run, result)
        binding = {"id":"p", "kind":"parameter", "source_id":self.ids["params"], "source_path":"/entries/0/current_value"}
        cases = [("[[source:p]]\n\nComputed error is 999 m. " + self.marked(claim), binding),
                 ("[[source:p]] Computed error is 999 m.\n\n" + self.marked(claim), binding),
                 ("[[source:p]]\n\n" + self.marked(claim), dict(binding, text="error is 999")),
                 ("[[source:p]]\n\n" + self.marked(claim), dict(binding, source_id=result, source_path="/metrics/value")),
                 ("[[source:p]]\n\n" + self.marked(claim), dict(binding, source_path="/entries/0/unit")),
                 ("[[source:p]]\n\n" + self.marked(claim), dict(binding, source_path="/entries/../current_value")),
                 ("[[source:p]]\n\n" + self.marked(claim), dict(binding, kind="year"))]
        for text, item in cases:
            with self.subTest(binding=item, text=text):
                with self.assertRaises(ValueError):
                    self.section(text, [claim], source_bindings=[item])

    def test_undeclared_duplicate_or_fenced_directives_are_rejected(self):
        run, result = self.checked()
        claim = self.claim(run, result)
        binding = {"id":"p", "kind":"parameter", "source_id":self.ids["params"], "source_path":"/entries/0/current_value"}
        cases = [("[[source:p]]", []), ("", [binding]),
                 ("[[source:p]]\n\n[[source:p]]", [binding]),
                 ("```text\n[[source:p]]\n```", [binding])]
        for prefix, bindings in cases:
            with self.subTest(prefix=prefix, bindings=bindings):
                with self.assertRaises(ValueError):
                    self.section(prefix + "\n\n" + self.marked(claim), [claim], source_bindings=bindings)

    def test_changed_artifact_and_rules_snapshot_invalidate_sections(self):
        run, result = self.checked()
        claim = self.claim(run, result)
        artifact = self.artifact()
        a = self.section("[[figure:1]]\n\n" + self.marked(claim), [claim],
                         structure=[{"kind":"figure", "number":1, "artifact_id":artifact}])
        delivery_fixture.DeliveryTests.create_pack(self)
        cp = self.rt.read()["copilot"]
        b = self.section("[[source:y]]\n\n" + self.marked(claim), [claim], source_bindings=[
            {"id":"y", "kind":"year", "source_id":cp["current"]["rules.lock"], "source_path":"/rules_year"}])
        self.write("artifacts/figure.txt", "changed")
        self.write("rules.txt", "changed")
        cp = self.rt.read()["copilot"]
        self.assertFalse(usable(self.root, cp, a, verified=True))
        self.assertFalse(usable(self.root, cp, b, verified=True))

    def test_exact_result_remains_verified_and_rounding_without_contract_rejected(self):
        run, result = self.checked(182.8125)
        self.claim(run, result, "Computed length is 182.8125 m.")
        with self.assertRaises(ValueError):
            self.claim(run, result, "Computed length is approximately 182.81 m.")
        with self.assertRaises(ValueError):
            self.claim(run, result, "Computed length is 182.8126 m.")

    def test_explicit_rounding_is_verified_without_rewriting_raw_metrics(self):
        run, result = self.checked(182.8125)
        before = copy.deepcopy(self.rt.read()["copilot"]["objects"][result])
        text = "Computed length is 182.81 m."
        contract = self.contract(result)
        claim = self.claim(run, result, text, [contract])
        self.section("# 1 模型结果\n\n" + self.marked(claim, text), [claim])
        cp = self.rt.read()["copilot"]
        self.assertEqual(cp["objects"][result], before)
        self.assertEqual(cp["objects"][claim]["payload"]["display_contracts"], [contract])
        self.rt.cover(self.rev(), "REQ-Q1-001", {"value":claim})

    def test_half_even_and_half_up_have_distinct_reproducible_ties(self):
        run, result = self.checked(2.125)
        for policy, shown in (("half_even", "2.12"), ("half_up", "2.13")):
            with self.subTest(policy=policy):
                contract = self.contract(result, "2.125", shown, rounding=policy)
                self.claim(run, result, "Computed length is " + shown + " m.", [contract])
                other = "2.13" if shown == "2.12" else "2.12"
                with self.assertRaises(ValueError):
                    self.claim(run, result, "Computed length is " + other + " m.", [dict(contract, display_value=other)])

    def test_percentage_contract_preserves_scale_and_rejects_decimal_confusion(self):
        run, result = self.checked(0.12345)
        contract = self.contract(result, "0.12345", "12.35%", format="percent", rounding="half_up")
        text = "Computed proportion is 12.35%."
        claim = self.claim(run, result, text, [contract])
        self.section(self.marked(claim, text), [claim])
        for shown in ("12.35", "0.12%", "12.34%"):
            with self.subTest(shown=shown):
                with self.assertRaises(ValueError):
                    self.claim(run, result, "Computed proportion is " + shown + ".", [dict(contract, display_value=shown)])

    def test_scientific_raw_value_uses_declared_fixed_precision(self):
        run, result = self.checked(0.0000001234)
        contract = self.contract(result, "1.234e-7", "0.00000012", decimal_places=8)
        self.claim(run, result, "Computed length is 0.00000012 m.", [contract])
        with self.assertRaises(ValueError):
            self.claim(run, result, "Computed length is 0.00 m.", [dict(contract, decimal_places=2, display_value="0.00")])

    def test_invalid_display_contract_fields_fail_closed_for_every_claim_label(self):
        run, result = self.checked(182.8125)
        good = self.contract(result)
        mutations = [{"raw_value":"999"}, {"display_value":"999.00"}, {"display_value":"182.82"},
            {"metric_path":"/missing"}, {"metric_path":"/value/../../value"}, {"metric_path":"/val~2ue"},
            {"result_id":self.ids["params"]}, {"version":"99"}, {"decimal_places":True},
            {"decimal_places":-1}, {"decimal_places":2.0}, {"decimal_places":100000},
            {"rounding":"floor"}, {"format":"expression"}, {"raw_value":"NaN"},
            {"raw_value":"Infinity"}, {"raw_value":182.8125}, {"verified":True}]
        for kind in ("numerical", "method", "comparison", "conclusion"):
            for mutation in mutations:
                with self.subTest(kind=kind, mutation=mutation):
                    before = self.rt.store.path.read_bytes()
                    with self.assertRaises(ValueError):
                        self.rt.claim(self.rev(), self.payload(run, "Computed length is 182.81 m.", [dict(good, **mutation)], kind), [result])
                    self.assertEqual(self.rt.store.path.read_bytes(), before)

    def test_unused_duplicate_and_missing_field_contracts_are_rejected(self):
        run, result = self.checked(182.8125)
        good = self.contract(result)
        missing = dict(good)
        del missing["rounding"]
        cases = [("Computed length is 182.8125 m.", [good]),
                 ("Computed length is 182.81 m.", [good, good]),
                 ("Computed length is 182.81 m.", [missing]),
                 ("Computed length is 182.81 m.", good)]
        for text, contracts in cases:
            with self.subTest(text=text, contracts=contracts):
                with self.assertRaises(ValueError):
                    self.rt.claim(self.rev(), self.payload(run, text, contracts), [result])

    def test_display_permission_does_not_spread_to_extra_values(self):
        run, result = self.checked(182.8125)
        contract = self.contract(result)
        text = "Computed length is 182.81 m."
        claim = self.claim(run, result, text, [contract])
        for extra in (" Error is 999 m.", " Range is 182.81–999 m.", " Total is 182.81+999 m."):
            with self.subTest(extra=extra):
                with self.assertRaises(ValueError):
                    self.claim(run, result, text + extra, [contract])
                with self.assertRaises(ValueError):
                    self.section(self.marked(claim, text + extra), [claim])

    def test_changed_parameters_invalidate_rounded_claim_and_section(self):
        run, result = self.checked(182.8125)
        text = "Computed length is 182.81 m."
        claim = self.claim(run, result, text, [self.contract(result)])
        section = self.section(self.marked(claim, text), [claim])
        self.params(100)
        cp = self.rt.read()["copilot"]
        for oid in (result, claim, section):
            with self.subTest(oid=oid):
                self.assertFalse(usable(self.root, cp, oid, verified=True))
        with self.assertRaises(ValueError):
            self.rt.claim(self.rev(), self.payload(run, text, [self.contract(result)]), [result])

    def test_physical_result_drift_prevents_display_source_reuse(self):
        run, result = self.checked(182.8125)
        obj = self.rt.read()["copilot"]["objects"][run]["payload"]
        self.write(obj["outputs"][0]["path"], {"value":999})
        with self.assertRaises(ValueError):
            self.rt.claim(self.rev(), self.payload(run, "Computed length is 182.81 m.", [self.contract(result)]), [result])

    def test_presealed_cached_bindings_cannot_validate_a_wrong_display_contract(self):
        run, result = self.checked(182.8125)
        payload = self.payload(run, "Computed length is 999 m.", [self.contract(result, shown="999")])
        payload.update(status="pass", content_assurance={"verified":True},
            numeric_bindings=[{"text":"999", "sources":[{"result_id":result, "value":999}]}])
        response = self.rt._tx(self.rev(), "qa", "Inject explicitly untrusted legacy fixture", lambda log:
            {"id":_put(log, "EvidenceMapEntry", "claim.legacy-display", payload, [result], status="verified")})
        claim = response["result"]["id"]
        cp = self.rt.read()["copilot"]
        self.assertFalse(usable(self.root, cp, claim, verified=True), object_errors(self.root, cp, claim))
        with self.assertRaises(ValueError):
            self.section(self.marked(claim, payload["claim"]), [claim])

    def test_negative_sign_is_not_lost_under_a_positive_display_contract(self):
        run, result = self.checked(182.8125)
        contract = self.contract(result)
        for text in ("Computed length is − 182.81 m.", "Computed length is -\n182.81 m."):
            with self.subTest(text=text):
                with self.assertRaises(ValueError):
                    self.claim(run, result, text, [contract])

    def test_rounding_can_reach_historical_audit_package_without_formal_ready(self):
        # Reuse the original delivery fixture's documented synthetic review data.
        # This checks plumbing, not a real human/render-engine verification.
        run, result = self.checked(182.8125)
        self.run = run
        self.claim_text = "Computed length is 182.81 m."
        claim = self.claim(run, result, self.claim_text, [self.contract(result)])
        self.rt.cover(self.rev(), "REQ-Q1-001", {"value":claim})
        section = self.section("# 1 模型结果\n\n" + self.marked(claim, self.claim_text), [claim])
        delivery_fixture.DeliveryTests.create_pack(self)
        delivery_fixture.DeliveryTests.create_paper(self)
        from docx import Document
        doc = Document(self.root / "paper.docx")
        heading = doc.add_heading("1 模型结果", level=1)
        doc._element.body.insert(0, heading._element)
        doc.save(self.root / "paper.docx")
        self.refresh_synthetic_paper_bindings()
        manifest = {"files":[{"path":"paper.docx", "role":"paper", "metadata_path":"paper.metadata.json", "visual_qa_path":"visual.json"}],
            "sections":[section], "claims":[claim],
            "paper_source":{"version":"0.1", "sections":[section]},
            "reviews":[{"kind":kind, "path":kind + ".json"} for kind in ("anonymity", "content")]}
        audit = self.delivery.audit(self.rev(), manifest)["result"]
        self.assertTrue(audit["passed"], audit)
        self.delivery.freeze(self.rev())
        package = self.delivery.package(self.rev(), "delivery/display.zip")["result"]
        self.assertTrue(audit_delivery_zip(self.root / package["path"])["passed"])
        self.assertFalse(submission_status(self.root, self.rt.read())["ready"])
        doc = Document(self.root / "paper.docx")
        doc.paragraphs[1].text = "Computed length is 999 m."
        doc.save(self.root / "paper.docx")
        self.refresh_synthetic_paper_bindings()
        changed = self.delivery.audit(self.rev(), manifest)["result"]
        self.assertFalse(changed["passed"], changed)
        self.assertFalse(changed["reports"]["paper_source"]["passed"], changed)


if __name__ == "__main__":
    unittest.main()
