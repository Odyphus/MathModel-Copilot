"""Real DOCX integration regressions found while auditing the historical paper."""
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import copilot_document_checks as adapter
from migrated.docx_audit import _reader_visible_math_source_issues


def bookmark(paragraph, name, number, text):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), name)
    paragraph._p.append(start)
    paragraph.add_run(text)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(number))
    paragraph._p.append(end)


class CitationAdapterTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.docx = self.root / "paper.docx"
        self.registry = self.root / "refs.json"
        # Synthetic local source, not a claim of external source verification.
        self.payload = {"schema_version":"1.0", "record_type":"reference_registry",
            "citation_style":{"marker_format":"[{number}]"},
            "references":[{"reference_id":"A", "number":1, "title":"Local test source",
                "bibliography_anchor":"BIB_A", "verification_status":"verified",
                "adoption_status":"adopted", "verification_record_hash":"1" * 64,
                "source_language":"en", "bibliography_language":"en",
                "foreign_source_reason":"Synthetic English fixture."}],
            "citations":[{"citation_id":"C", "reference_id":"A", "claim_id":"CLAIM-A@1",
                "number":1, "anchor":"CITE_A"}]}
        document = Document()
        bookmark(document.add_paragraph("Source "), "CITE_A", 1, "[1]")
        bookmark(document.add_paragraph(), "BIB_A", 2, "[1] Local test source")
        document.save(self.docx)
        self.write_registry()

    def write_registry(self):
        self.registry.write_text(json.dumps(self.payload), encoding="utf-8")

    def test_real_bidirectional_pass_reaches_adapter_boolean(self):
        result = adapter.audit_citations(self.docx, self.registry)
        self.assertEqual(result["hard_gate_status"], "PASS")
        self.assertEqual(result["migrated_status"], "READY")
        self.assertTrue(result["passed"])
        self.assertEqual(result["status"], "pass")

    def test_real_missing_bookmark_is_not_promoted(self):
        document = Document()
        document.add_paragraph("[1] Local test source")
        document.save(self.docx)
        result = adapter.audit_citations(self.docx, self.registry)
        self.assertFalse(result["passed"])
        self.assertEqual(result["hard_gate_status"], "FAIL")

    def test_real_unverified_reference_is_not_promoted(self):
        self.payload["references"][0]["verification_status"] = "generated"
        self.write_registry()
        self.assertFalse(adapter.audit_citations(self.docx, self.registry)["passed"])

    def test_warning_is_preserved_without_overriding_hard_gate(self):
        result = adapter.audit_citations(self.docx, self.registry)
        self.assertTrue(result["passed"])
        self.assertEqual(result["quality_status"], "WARN")
        self.assertTrue(any(x["severity"] == "warning" and x["status"] == "fail" for x in result["findings"]))

    def test_incomplete_or_contradictory_checker_contract_fails_closed(self):
        for report in ({"passed": True}, {"hard_gate_status":"PASS", "status":"READY", "findings":[]},
                {"hard_gate_status":"PASS", "status":"READY", "findings":[{"severity":"error", "status":"fail"}]}):
            with self.subTest(report=report), patch("migrated.citation_audit.audit_citations", return_value=dict(report)):
                self.assertFalse(adapter.audit_citations(self.docx, self.registry)["passed"])


class BoundFileLabelTests(unittest.TestCase):
    path = "assets/Case_1_ result_E.xls"

    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.docx = Path(temporary.name) / "paper.docx"

    def make_docx(self, text, *, appendix=True):
        document = Document()
        if appendix:
            document.add_paragraph("附录")
        document.add_paragraph(text)
        document.save(self.docx)

    def test_actual_spaced_file_label_needs_exact_source_binding(self):
        self.make_docx("文件：" + self.path)
        self.assertTrue(_reader_visible_math_source_issues(self.docx))
        self.assertEqual(_reader_visible_math_source_issues(self.docx, file_reference_paths=(self.path,)), [])

    def test_bound_path_cannot_hide_equation_in_same_paragraph(self):
        for suffix in (" and x_i", " + y_i", "^2", " and 999 m"):
            with self.subTest(suffix=suffix):
                self.make_docx("文件：" + self.path + suffix)
                self.assertTrue(_reader_visible_math_source_issues(self.docx, file_reference_paths=(self.path,)))

    def test_different_or_prefix_path_is_not_authorized(self):
        self.make_docx("文件：" + self.path)
        for binding in ("assets/", "assets/Case_2_ result_E.xls", "Case_1_ result_E.xls"):
            with self.subTest(binding=binding):
                self.assertTrue(_reader_visible_math_source_issues(self.docx, file_reference_paths=(binding,)))

    def test_same_label_in_body_is_still_checked(self):
        self.make_docx("文件：" + self.path, appendix=False)
        self.assertTrue(_reader_visible_math_source_issues(self.docx, file_reference_paths=(self.path,)))

    def test_prose_prefix_does_not_turn_into_file_list_block(self):
        self.make_docx("Example 文件：" + self.path)
        self.assertTrue(_reader_visible_math_source_issues(self.docx, file_reference_paths=(self.path,)))

    def test_adjacent_unbuilt_formula_remains_a_failure(self):
        document = Document()
        document.add_paragraph("附录")
        document.add_paragraph("文件：" + self.path)
        document.add_paragraph("The value is x_i.")
        document.save(self.docx)
        issues = _reader_visible_math_source_issues(self.docx, file_reference_paths=(self.path,))
        self.assertEqual(len(issues), 1)
        self.assertIn("x_i", issues[0])

    def test_public_document_adapter_retains_exact_path_scope(self):
        self.make_docx("文件：" + self.path)
        policy = {"profile":"cumcm", "require_visual_qa":False}
        without = adapter.audit_docx(self.docx, policy=policy)
        with_proof = adapter.audit_docx(self.docx, policy=policy, file_reference_paths=(self.path,))
        check = lambda report: next(x["status"] for x in report["findings"] if x["check_id"] == "docx.no_reader_visible_math_source")
        self.assertEqual(check(without), "fail")
        self.assertEqual(check(with_proof), "pass")
        # Other actual document errors are not overridden by this file-name fix.
        self.assertFalse(with_proof["passed"])


if __name__ == "__main__":
    unittest.main()
