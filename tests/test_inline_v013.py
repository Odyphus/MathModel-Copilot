"""Source-consumption regressions for explicit DOCX inline support."""
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "scripts"), str(ROOT / "tests")]
import test_paper_source_v012 as fixtures
from copilot_paper_source import read_docx_blocks
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree


class InlineV013Tests(unittest.TestCase):
    def setUp(self):
        self.paper = fixtures.PaperSourceV012Tests()
        self.paper.setUp()
        self.addCleanup(self.paper.doCleanups)

    def symbol(self, run, text="999", font="Arial"):
        for char in text:
            node = OxmlElement("w:sym")
            node.set(qn("w:font"), font)
            node.set(qn("w:char"), format(ord(char), "04X"))
            run._r.append(node)

    def assert_rejected(self, doc, reason):
        path = self.paper.save(doc)
        result = self.paper.check(path)
        self.assertFalse(result["passed"], result)
        self.assertIn(reason, " ".join(result["errors"]))
        return path

    def test_font_encoded_symbols_after_real_claim_are_explicitly_rejected(self):
        for value, font in (("999", "Arial"), ("7", "Calibri"), ("-", "Symbol"), ("α", "Wingdings")):
            with self.subTest(value=value, font=font):
                doc = Document(self.paper.root / "paper.docx")
                self.symbol(doc.paragraphs[0].add_run(), value, font)
                self.assert_rejected(doc, "sym")

    def test_symbol_only_paragraph_is_not_an_empty_paragraph(self):
        doc = Document(self.paper.root / "paper.docx")
        self.symbol(doc.add_paragraph().add_run(), "58")
        path = self.assert_rejected(doc, "sym")
        with self.assertRaisesRegex(ValueError, "sym"):
            read_docx_blocks(path)

    def test_symbols_in_table_cell_use_the_same_rejection(self):
        self.paper.section_text("| Finding |\n| --- |\n| " + self.paper.paragraph() + " |")
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "Finding"
        table.cell(1, 0).text = self.paper.text
        self.symbol(table.cell(1, 0).paragraphs[0].add_run(), "123")
        self.assert_rejected(doc, "sym")

    def test_symbols_in_bound_code_appendix_cannot_disappear(self):
        from docx.enum.style import WD_STYLE_TYPE
        self.paper.contract["supplements"] = [{"kind":"source_code",
            "source_id":self.paper.fixture.ids["code"], "path":"solver.py"}]
        doc = Document(self.paper.root / "paper.docx")
        doc.styles.add_style("CopilotSourceCode", WD_STYLE_TYPE.PARAGRAPH)
        doc.add_paragraph("源程序：solver.py")
        p = doc.add_paragraph((self.paper.root / "solver.py").read_text(encoding="utf-8").rstrip("\n"), "CopilotSourceCode")
        self.assertTrue(self.paper.check(self.paper.save(doc))["passed"])
        self.symbol(p.add_run(), "5")
        self.assert_rejected(doc, "sym")

    def test_unknown_run_children_are_not_treated_as_empty(self):
        for name in ("w:unknownVisible", "w:object", "w:pict", "w:footnoteRef"):
            with self.subTest(name=name):
                doc = Document(self.paper.root / "paper.docx")
                doc.paragraphs[0].add_run()._r.append(OxmlElement(name))
                self.assert_rejected(doc, name.split(":")[1])

    def test_unknown_namespace_is_not_accepted_by_local_name(self):
        doc = Document(self.paper.root / "paper.docx")
        doc.paragraphs[0].add_run()._r.append(etree.Element("{urn:unknown-visible}t", char="0039"))
        self.assert_rejected(doc, "urn:unknown-visible")

    def test_unknown_paragraph_wrapper_cannot_hide_behind_descendant_text(self):
        doc = Document(self.paper.root / "paper.docx")
        p = doc.paragraphs[0]._p
        run = doc.paragraphs[0].runs[0]._r
        p.remove(run)
        wrapper = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        content.append(run); wrapper.append(content); p.append(wrapper)
        self.assert_rejected(doc, "sdt")

    def test_page_break_cannot_join_two_numeric_tokens(self):
        doc = Document(self.paper.root / "paper.docx")
        p = doc.paragraphs[0]
        p.text = "Computed length is 1"
        br = OxmlElement("w:br"); br.set(qn("w:type"), "page")
        p.add_run()._r.append(br)
        p.add_run("0 m.")
        self.assertFalse(self.paper.check(self.paper.save(doc))["passed"])

    def test_supported_tabs_breaks_and_hyphens_preserve_visible_text(self):
        self.paper.section_text("A co-operative model.\n\n" + self.paper.paragraph())
        doc = Document()
        p = doc.add_paragraph("A\tco")
        p.add_run()._r.append(OxmlElement("w:noBreakHyphen"))
        p.add_run("operative")
        p.add_run()._r.append(OxmlElement("w:cr"))
        p.add_run("model.")
        doc.add_paragraph(self.paper.text)
        report = self.paper.check(self.paper.save(doc))
        self.assertTrue(report["passed"], report)

    def test_supported_image_rejects_additional_drawing_text(self):
        from PIL import Image
        path = self.paper.root / "figure.png"
        Image.new("RGB", (20, 20), "navy").save(path)
        result_id = next(o["id"] for o in self.paper.rt.read()["copilot"]["objects"].values() if o["kind"] == "ResultRecord")
        artifact = self.paper.rt.register(self.paper.fixture.rev(), "ArtifactRecord", "figure.bound",
            {"artifact_id":"figure.bound", "artifact_type":"figure", "path":"figure.png"},
            dependencies=[result_id], files=["figure.png"])["result"]["object_id"]
        self.paper.section_text("[[figure:1]]\n\n![Source](figure.png)\n\n" + self.paper.paragraph(),
            structure=[{"kind":"figure", "number":1, "artifact_id":artifact}])
        doc = Document(); doc.add_paragraph("图 1"); doc.add_picture(str(path)); doc.add_paragraph(self.paper.text)
        self.assertTrue(self.paper.check(self.paper.save(doc))["passed"])
        drawing = next(doc.element.iter(qn("w:drawing")))
        extra = OxmlElement("a:t"); extra.text = "999"
        drawing.append(extra)
        self.assert_rejected(doc, "drawing")

    def test_symbol_in_native_math_run_is_not_ignored(self):
        self.paper.section_text("$$x$$\n\n" + self.paper.paragraph())
        doc = Document(); p = doc.add_paragraph()
        math, run, text = OxmlElement("m:oMath"), OxmlElement("m:r"), OxmlElement("m:t")
        text.text = "x"; run.append(text); math.append(run); p._p.append(math)
        doc.add_paragraph(self.paper.text)
        self.assertTrue(self.paper.check(self.paper.save(doc))["passed"])
        sym = OxmlElement("w:sym"); sym.set(qn("w:font"), "Arial"); sym.set(qn("w:char"), "0039")
        run.append(sym)
        self.assert_rejected(doc, "sym")


if __name__ == "__main__":
    unittest.main()
