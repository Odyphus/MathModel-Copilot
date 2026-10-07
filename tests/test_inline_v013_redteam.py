"""Independent adjacent DOCX probes; assertions cover source audit only.

The documents are synthetic and are created from real registered fixture
objects. Tests never fabricate verified objects or modify the authority JSON.
"""
from __future__ import annotations

import hashlib
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "scripts"), str(ROOT / "tests")]

import test_paper_source_v012 as fixtures
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


class InlineV013RedTeamTests(unittest.TestCase):
    def setUp(self):
        self.fixture = fixtures.PaperSourceV012Tests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)

    def doc(self):
        return Document(self.fixture.root / "paper.docx")

    def check(self, document):
        path = self.fixture.save(document)
        authority = self.fixture.root / "state/decision_log.json"
        before = hashlib.sha256(authority.read_bytes()).hexdigest()
        report = self.fixture.check(path)
        self.assertEqual(before, hashlib.sha256(authority.read_bytes()).hexdigest())
        return report

    def reject(self, document):
        report = self.check(document)
        self.assertFalse(report["passed"], report)
        self.assertTrue(report["errors"], report)

    def table_doc(self):
        f = self.fixture
        f.section_text("| Finding |\n| --- |\n| " + f.paragraph() + " |")
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "Finding"
        table.cell(1, 0).text = f.text
        return doc, table.cell(1, 0).paragraphs[0]

    def sign_before_result(self, paragraph):
        left, right = paragraph.text.split("10", 1)
        paragraph.clear()
        paragraph.add_run(left)._r.append(OxmlElement("w:noBreakHyphen"))
        paragraph.add_run("10" + right)

    def split_result(self, paragraph):
        left, right = paragraph.text.split("10", 1)
        paragraph.clear()
        paragraph.add_run(left + "1")._r.append(OxmlElement("w:cr"))
        paragraph.add_run("0" + right)

    def test_nonbreaking_sign_cannot_disappear_before_result(self):
        doc = self.doc()
        self.sign_before_result(doc.paragraphs[0])
        self.reject(doc)

    def test_nonbreaking_sign_cannot_disappear_in_table_result(self):
        doc, paragraph = self.table_doc()
        self.sign_before_result(paragraph)
        self.reject(doc)

    def test_carriage_return_cannot_join_distinct_numeric_tokens(self):
        doc = self.doc()
        self.split_result(doc.paragraphs[0])
        self.reject(doc)

    def test_carriage_return_cannot_join_table_numeric_tokens(self):
        doc, paragraph = self.table_doc()
        self.split_result(paragraph)
        self.reject(doc)

    def test_uncached_dynamic_field_is_not_a_static_source(self):
        doc = self.doc()
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "= 777")
        doc.paragraphs[0]._p.append(field)
        self.reject(doc)

    def test_dynamic_field_in_table_is_not_silently_ignored(self):
        doc, paragraph = self.table_doc()
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "DATE")
        paragraph._p.append(field)
        self.reject(doc)

    def test_conditional_soft_hyphen_cannot_be_assumed_invisible(self):
        doc = self.doc()
        doc.paragraphs[0].add_run()._r.append(OxmlElement("w:softHyphen"))
        self.reject(doc)

    def test_unknown_inline_wrapper_cannot_authorize_registered_text(self):
        doc = self.doc()
        paragraph = doc.paragraphs[0]._p
        # No font/rendering assumption is made about a foreign namespace.
        # Unknown behavior must be explicit, not recursively flattened.
        from lxml import etree
        wrapper = etree.Element("{urn:copilot-redteam:unknown-rendering}content")
        for run in list(paragraph):
            paragraph.remove(run)
            wrapper.append(run)
        paragraph.append(wrapper)
        self.reject(doc)

    def test_word_character_in_omml_run_cannot_disappear(self):
        f = self.fixture
        f.section_text("$$x$$\n\n" + f.paragraph())
        doc = Document()
        paragraph = doc.add_paragraph()
        equation, run, text = OxmlElement("m:oMath"), OxmlElement("m:r"), OxmlElement("m:t")
        text.text = "x"
        run.append(text)
        run.append(OxmlElement("w:noBreakHyphen"))
        equation.append(run)
        paragraph._p.append(equation)
        doc.add_paragraph(f.text)
        self.reject(doc)

    def code_doc(self):
        f = self.fixture
        f.contract["supplements"] = [{"kind": "source_code", "source_id": f.fixture.ids["code"], "path": "solver.py"}]
        doc = self.doc()
        doc.styles.add_style("CopilotSourceCode", WD_STYLE_TYPE.PARAGRAPH)
        doc.add_paragraph("源程序：solver.py")
        code = doc.add_paragraph((f.root / "solver.py").read_text(encoding="utf-8").rstrip("\n"), "CopilotSourceCode")
        return doc, code

    def test_supported_carriage_return_preserves_exact_code_lines(self):
        doc, code = self.code_doc()
        breaks = code._p.findall(".//" + qn("w:br"))
        self.assertTrue(breaks, "fixture must contain actual source-code line breaks")
        for node in breaks:
            node.tag = qn("w:cr")
        report = self.check(doc)
        self.assertTrue(report["passed"], report)

    def test_extra_carriage_return_in_code_is_not_normalized_away(self):
        doc, code = self.code_doc()
        run = code.runs[0]._r
        first = run.find(qn("w:t"))
        self.assertIsNotNone(first)
        value = first.text
        self.assertGreater(len(value), 3)
        first.text = value[:2]
        index = list(run).index(first)
        suffix = OxmlElement("w:t")
        suffix.text = value[2:]
        run.insert(index + 1, OxmlElement("w:cr"))
        run.insert(index + 2, suffix)
        self.reject(doc)

    def test_nonvisible_bookmark_and_proofing_markers_preserve_source(self):
        doc = self.doc()
        paragraph = doc.paragraphs[0]._p
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), "0")
        start.set(qn("w:name"), "source_finding")
        end.set(qn("w:id"), "0")
        spell_start, spell_end = OxmlElement("w:proofErr"), OxmlElement("w:proofErr")
        spell_start.set(qn("w:type"), "spellStart")
        spell_end.set(qn("w:type"), "spellEnd")
        paragraph.insert(0, start)
        paragraph.insert(1, spell_start)
        paragraph.append(spell_end)
        paragraph.append(end)
        report = self.check(doc)
        self.assertTrue(report["passed"], report)

    def test_hyperlink_label_is_read_in_order(self):
        doc = self.doc()
        paragraph = doc.paragraphs[0]._p
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), "source_finding")
        for run in list(paragraph):
            paragraph.remove(run)
            link.append(run)
        paragraph.append(link)
        report = self.check(doc)
        self.assertTrue(report["passed"], report)
        text = link.find(".//" + qn("w:t"))
        text.text = text.text.replace("10", "777")
        self.reject(doc)


if __name__ == "__main__":
    unittest.main()
