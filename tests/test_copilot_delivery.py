"""Synthetic fixtures exercise delivery; they are not real human signoffs."""
from __future__ import annotations

import copy
import json
import struct
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))
from test_copilot_runtime import make_project
import copilot_domain as domain
from copilot_delivery import Delivery, submission_status, audit_delivery_zip, MANIFEST_ARCHIVE_NAME, _ai_checks
from copilot_packs import load_pack
from copilot_runtime import bind_file


class DeliveryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.rt, self.ids, self.reg = make_project(self.root, count=1)
        self.delivery = Delivery(self.root)
        self.claim_text = "Computed length is 10 m."
        self.run = self.rt.execute(self.rev(), "Q1", ["{python}", "solver.py"], ["result.json"],
            dependencies=[self.ids[k] for k in ("model", "params", "data", "code", "plan")])["result"]["object_id"]
        check = self.rt.validate_run(self.rev(), self.run, "checker.py", ["{python}", "{checker}", "{run}", "{report}"])["result"]
        self.assertTrue(check["passed"], check)
        run = self.rt.read()["copilot"]["objects"][self.run]["payload"]
        self.claim = self.rt.claim(self.rev(), {
            "claim_id": "C1", "claim": self.claim_text, "claim_type": "numerical", "paper_anchor": "results",
            "limitations": ["Only this deterministic synthetic test fixture."], "formal_run_id": self.run,
            "requirement_ids": ["REQ-Q1-001"], "data_sources": [run["data_hash"]],
            "code_locations": ["solver.py"], "tables": [run["outputs"][0]["path"]],
            "validation_evidence": ["Q1.run.formal"],
        }, [check["result_id"]])["result"]["claim_id"]
        self.rt.cover(self.rev(), "REQ-Q1-001", {"value": self.claim})
        self.write("section.md", self.claim_text + " [[claim:" + self.claim + "]]\n")
        self.section = self.delivery.section(self.rev(), "paper.results", "section.md", [self.claim])["result"]["section_id"]
        self.create_pack()
        self.create_paper()
        self.manifest = {
            "files": [{"path": "paper.docx", "role": "paper", "metadata_path": "paper.metadata.json", "visual_qa_path": "visual.json"}],
            "sections": [self.section], "claims": [self.claim],
            "reviews": [{"kind": kind, "path": kind + ".json"} for kind in ("anonymity", "content")],
        }

    def rev(self):
        return self.rt.read()["copilot"]["revision"]

    def write(self, path, value):
        target = self.root / path
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(value if isinstance(value, str) else json.dumps(value), encoding="utf-8")

    def create_pack(self):
        pack = {k: v for k, v in load_pack("generic").items() if not k.startswith("_")}
        pack.update(id="fixture-contest", name="Synthetic delivery fixture", aliases=[], rules_year=2026,
            rules_status="baseline", resources={}, document_profile="generic",
            official_sources=[{"id": "fixture", "url": "https://example.edu/fixture/rules-2026",
                "origin": "official_rule", "rules_year": 2026, "required_for_lock": True, "checked_at": "2026-10-05"}],
            rules={key: {"value": value, "origin": "official_rule", "source_ids": ["fixture"]}
                for key, value in {"paper_formats": ["docx"], "anonymity": True, "separate_support_archive": False}.items()})
        self.write("competition/pack.json", pack)
        self.write("rules.txt", "Synthetic rules for a unit test; not an actual contest source.")
        # The path form also exercises the CLI's string-pack interface.
        self.delivery.lock_rules(self.rev(), "competition/pack.json", 2018, 2026, "historical_benchmark",
            [{"url": pack["official_sources"][0]["url"], "path": "rules.txt", "sha256": domain.sha256_file(self.root / "rules.txt"),
              "verified_at": "2026-10-05", "rules_year": 2026}], "qa-fixture", "integrator")

    def relock(self, changes=None, formal=False):
        pack_path = self.root / "competition/pack.json"
        pack = json.loads(pack_path.read_text(encoding="utf-8"))
        if changes:
            changes(pack)
        self.write("competition/pack.json", pack)
        if formal:
            self.rt.configure(self.rev(), problem_year=2026, rules_year=2026, evaluation_mode="formal_contest")
        self.delivery.lock_rules(self.rev(), str(pack_path), 2026 if formal else 2018, 2026,
            "formal_contest" if formal else "historical_benchmark", [{"url": pack["official_sources"][0]["url"],
            "path": "rules.txt", "sha256": domain.sha256_file(self.root / "rules.txt"), "verified_at": "2026-10-05"}], "qa-fixture")

    def create_paper(self, extra=""):
        from docx import Document
        from PIL import Image
        doc = Document()
        doc.add_paragraph(self.claim_text)
        if extra:
            doc.add_paragraph(extra)
        doc.core_properties.author = ""
        doc.core_properties.last_modified_by = ""
        doc.save(self.root / "paper.docx")
        sha = domain.sha256_file(self.root / "paper.docx")
        self.write("paper.metadata.json", {"sha256": sha.lower(), "milestone": "working", "history_class": "working_build",
            "formal_history_eligible": False, "artifact_id": "ART-fixture", "formal_values_present": True, "run_id": self.run})
        image = self.root / "render/page1.png"
        image.parent.mkdir(exist_ok=True)
        Image.new("RGB", (20, 20), "white").save(image)
        self.write("visual.json", {"source_docx_sha256": sha.lower(), "status": "pass", "page_count": 1, "inspected_pages": [1],
            "actor_kind": "agent_evaluator", "reviewer": "qa-fixture", "notes": ["Synthetic image and inspection fixture; no human review asserted."],
            "rendered_pages": [{"page": 1, "path": "render/page1.png", "sha256": domain.sha256_file(image)}]})
        for kind in ("anonymity", "content"):
            self.write(kind + ".json", {"kind": kind, "result": "pass", "actor_kind": "agent_evaluator", "reviewer": "qa-fixture",
                "reviewed_at": "2026-10-05", "notes": ["Synthetic inspection evidence for tests only."],
                "files": [{"path": "paper.docx", "sha256": sha}]})

    def audit(self):
        return self.delivery.audit(self.rev(), self.manifest)["result"]

    def package(self):
        audit = self.audit()
        self.assertTrue(audit["passed"], audit)
        self.delivery.freeze(self.rev())
        return self.delivery.package(self.rev(), "delivery/package.zip")["result"]

    def test_historical_checked_frozen_package_is_not_formal_ready(self):
        result = self.package()
        log = self.rt.read()
        status = submission_status(self.root, log)
        self.assertEqual(status["state"], "frozen")
        self.assertFalse(status["ready"])
        self.assertTrue(any("Historical" in error for error in status["blockers"]))
        report = audit_delivery_zip(self.root / result["path"])
        self.assertTrue(report["passed"], report)
        with zipfile.ZipFile(self.root / result["path"]) as archive:
            self.assertEqual(set(archive.namelist()), {"paper.docx", MANIFEST_ARCHIVE_NAME})
            manifest = json.loads(archive.read(MANIFEST_ARCHIVE_NAME))
            self.assertEqual(manifest["files"][0]["sha256"], domain.sha256_file(self.root / "paper.docx"))
            self.assertIsNone(archive.testzip())
        obj = log["copilot"]["objects"][result["package_id"]]
        self.assertTrue(obj["payload"]["crc_checked"])

    def test_missing_question_fails_and_retains_audit_evidence(self):
        self.rt.configure(self.rev(), question_count=2)
        result = self.audit()
        self.assertFalse(result["passed"])
        self.assertEqual(result["reports"]["requirements"]["questions"]["Q2"], "missing")
        self.assertEqual(self.rt.read()["copilot"]["objects"][result["audit_id"]]["status"], "failed")
        with self.assertRaises(ValueError):
            self.delivery.freeze(self.rev())

    def test_section_rejects_missing_unverified_marker_and_unsupported_numbers(self):
        cases = [([], self.claim_text), (["missing@1"], self.claim_text + " [[claim:missing@1]]"),
            ([self.claim], self.claim_text), ([self.claim], "We obtained 99 m. [[claim:" + self.claim + "]]"),
            ([self.claim], self.claim_text + " Additional error is 99. [[claim:" + self.claim + "]]")]
        for claims, text in cases:
            with self.subTest(text=text):
                self.write("bad-section.md", text)
                with self.assertRaises(ValueError):
                    self.delivery.section(self.rev(), "paper.bad", "bad-section.md", claims)

    def test_final_paper_omitting_claim_cannot_pass(self):
        old = self.claim_text
        self.claim_text = "There is no reported result."
        self.create_paper()
        self.claim_text = old
        result = self.audit()
        self.assertFalse(result["passed"])
        self.assertTrue(any("omits evidence-bound" in e for e in result["errors"]), result)

    def test_boolean_reviews_cannot_pass(self):
        self.write("anonymity.json", {"kind": "anonymity", "result": "pass", "verified": True})
        result = self.audit()
        self.assertFalse(result["passed"])
        self.assertTrue(any("review" in e for e in result["errors"]), result)

    def test_real_anonymity_failure_overrides_positive_review_record(self):
        self.create_paper("Contact test@example.com for details.")
        result = self.audit()
        self.assertFalse(result["passed"])
        findings = result["reports"]["paper"]["findings"]
        self.assertTrue(any(f["check_id"] == "docx.privacy" and f["status"] == "fail" for f in findings), result)

    def test_paper_or_review_drift_invalidates_old_audit(self):
        result = self.audit()
        self.assertTrue(result["passed"], result)
        self.write("content.json", "{}")
        with self.assertRaises(ValueError):
            self.delivery.freeze(self.rev())
        self.assertFalse(submission_status(self.root, self.rt.read())["ready"])

    def test_changed_section_revision_invalidates_audit(self):
        result = self.audit()
        self.assertTrue(result["passed"], result)
        self.write("section.md", "# Revised\n\n" + self.claim_text + " [[claim:" + self.claim + "]]\n")
        self.delivery.section(self.rev(), "paper.results", "section.md", [self.claim])
        with self.assertRaises(ValueError):
            self.delivery.freeze(self.rev())

    def test_ai_usage_change_invalidates_frozen_audit(self):
        result = self.audit()
        self.assertTrue(result["passed"], result)
        self.delivery.freeze(self.rev())
        self.rt.log_ai(self.rev(), {"entry": {"tool": "fixture", "purpose": "synthetic test record"}})
        with self.assertRaisesRegex(ValueError, "basis changed"):
            self.delivery.package(self.rev(), "delivery/stale.zip")

    def test_numeric_spellings_do_not_reuse_different_claim_numbers(self):
        for value in (".10", "−10", "10e2", "10%"):
            with self.subTest(value=value):
                self.write("bad-section.md", self.claim_text + " Extra result is " + value + ". [[claim:" + self.claim + "]]")
                with self.assertRaisesRegex(ValueError, "Unsupported numerical"):
                    self.delivery.section(self.rev(), "paper.bad", "bad-section.md", [self.claim])

    def test_rules_snapshot_and_pack_drift_invalidate_frozen_result(self):
        self.package()
        self.write("rules.txt", "Changed rules snapshot")
        status = submission_status(self.root, self.rt.read())
        self.assertFalse(status["ready"])
        with self.assertRaises(ValueError):
            self.delivery.package(self.rev(), "delivery/second.zip")

    def test_pack_bytes_are_rechecked_even_without_object_dependency(self):
        result = self.audit()
        self.assertTrue(result["passed"], result)
        self.write("competition/pack.json", {})
        with self.assertRaises(ValueError):
            self.delivery.freeze(self.rev())

    def test_zip_change_invalidates_package(self):
        result = self.package()
        with zipfile.ZipFile(self.root / result["path"], "a") as archive:
            archive.writestr("unlisted.txt", "not in whitelist")
        status = submission_status(self.root, self.rt.read())
        self.assertTrue(any("package" in error.lower() for error in status["blockers"]), status)
        self.assertFalse(audit_delivery_zip(self.root / result["path"])["passed"])

    def test_paths_and_existing_assets_are_never_overwritten(self):
        original = self.package()
        before = (self.root / original["path"]).read_bytes()
        for target in ("../escape.zip", "C:/escape.zip", original["path"]):
            with self.subTest(target=target), self.assertRaises(ValueError):
                self.delivery.package(self.rev(), target)
        self.assertEqual(before, (self.root / original["path"]).read_bytes())

    def test_crc_failure_is_detected_on_actual_archive_bytes(self):
        result = self.package()
        source = self.root / result["path"]
        bad = self.root / "crc-bad.zip"
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(bad, "w", compression=zipfile.ZIP_STORED) as archive:
            for info in original.infolist():
                archive.writestr(info.filename, original.read(info.filename))
        with zipfile.ZipFile(bad) as archive:
            info = archive.getinfo("paper.docx")
        data = bytearray(bad.read_bytes())
        filename_len, extra_len = struct.unpack_from("<HH", data, info.header_offset + 26)
        index = info.header_offset + 30 + filename_len + extra_len + info.file_size // 2
        data[index] ^= 1
        bad.write_bytes(data)
        report = audit_delivery_zip(bad)
        self.assertFalse(report["passed"], report)
        self.assertTrue(any("CRC" in error for error in report["errors"]), report)

    def test_no_automatic_submission_and_positional_event_contract(self):
        self.package()
        self.assertEqual(self.rt.read()["copilot"]["delivery"].get("events", []), [])
        for state in ("awaiting_submission", "submitted", "receipt_received"):
            with self.subTest(state=state), self.assertRaises(ValueError):
                self.delivery.event(self.rev(), state, None, "agent", "qa-fixture")
        self.assertEqual(self.rt.read()["copilot"]["delivery"]["state"], "frozen")

    def test_unknown_document_formats_and_malformed_visual_records_fail_closed(self):
        self.write("visual.json", [])
        result = self.audit()
        self.assertFalse(result["passed"])
        self.assertEqual(self.rt.read()["copilot"]["objects"][result["audit_id"]]["status"], "failed")

    def test_receipt_transition_checks_previous_event_before_writing(self):
        from unittest.mock import patch
        from copilot_runtime import _put
        # Isolate the event-chain gate; this does not create a real formal audit.
        def fixture(log):
            log["problem_meta"]["evaluation_mode"] = "formal_contest"
            event = _put(log, "DeliveryEvent", "fixture.event", {"state":"submitted"}, status="stale")
            log["copilot"]["delivery"].update(state="submitted", events=[event])
        self.rt._tx(self.rev(), "synthetic-fixture", "Exercise stale event gate", fixture)
        before = self.rt.store.path.read_bytes()
        with patch("copilot_delivery._audit_blockers", return_value=[]), patch("copilot_delivery._package_errors", return_value=[]):
            with self.assertRaisesRegex(ValueError, "Previous delivery event"):
                self.delivery.event(self.rev(), "receipt_received", None, "human", "synthetic-fixture")
        self.assertEqual(before, self.rt.store.path.read_bytes())

    def test_unknown_required_rule_cannot_be_silently_ignored(self):
        self.relock(lambda p: p["required_checks"].append("fixture_unimplemented_rule"))
        result = self.audit()
        self.assertFalse(result["passed"])
        self.assertTrue(any("no implemented evaluator" in e for e in result["errors"]), result)

    def test_strict_pack_byte_limit_is_not_treated_as_inclusive(self):
        size = (self.root / "paper.docx").stat().st_size
        def limit(pack):
            pack["rules"]["paper_max_bytes"] = {"value": size, "comparison": "lt", "origin": "official_rule", "source_ids": ["fixture"]}
        self.relock(limit)
        result = self.audit()
        self.assertFalse(result["passed"])
        self.assertTrue(any("violates lt" in e for e in result["errors"]), result)

    def test_required_supporting_archive_cannot_be_omitted(self):
        def required(pack):
            pack["rules"]["separate_support_archive"]["value"] = True
            pack["required_checks"].append("supporting_materials_passed")
        self.relock(required)
        result = self.audit()
        self.assertFalse(result["passed"])
        self.assertTrue(any("actual supporting archive" in e for e in result["errors"]), result)

    def test_formal_mode_without_render_receipt_cannot_be_ready(self):
        self.relock(formal=True)
        self.rt.log_ai(self.rev(), {"action": "declare_none", "reason": "Synthetic no-use fixture, not a real team declaration."})
        for filename in ("visual.json", "anonymity.json", "content.json"):
            proof = json.loads((self.root / filename).read_text(encoding="utf-8"))
            proof["actor_kind"] = "human"  # Fixture exercises the declared-role gate only.
            self.write(filename, proof)
        result = self.audit()
        self.assertFalse(result["passed"])
        self.assertTrue(any("render receipt" in e for e in result["errors"]), result)
        self.assertTrue(any("References need" in e for e in result["errors"]), result)
        self.assertFalse(submission_status(self.root, self.rt.read())["ready"])

    def test_unknown_ai_ledger_does_not_count_as_no_use(self):
        def required(pack):
            pack["required_checks"].append("ai_disclosure_passed")
        self.relock(required)
        result = self.audit()
        self.assertFalse(result["passed"])
        self.assertFalse(result["reports"]["ai_disclosure"]["passed"])

    def test_mcm_ai_report_requires_real_trailing_page_interval(self):
        entry = {"tool": "TestTool", "model": "TestModel", "version": "test-v", "use_stage": "writing", "purpose": "fixture",
            "paper_sections": ["results"], "query": "test question", "output": "test response", "human_review": "independently checked fixture"}
        log = {"compliance": {"ai_usage": [entry]}}
        pack = {"id": "mcm", "rules": {"ai_report_position": {"value": "after_solution"}}, "required_checks": ["ai_disclosure_passed"]}
        report, _ = _ai_checks(self.root, log, pack, {}, "", ["solution", "Report on Use of AI"])
        self.assertFalse(report["passed"])
        self.assertTrue(any("page interval" in e for e in report["errors"]))

    def test_cumcm_nonempty_ledger_requires_details_in_actual_support_zip(self):
        from render_ai_usage import validate_entries, render_cumcm_use_statement
        entry = {"tool": "TestTool", "model": "TestModel", "version": "test-v", "use_stage": "writing", "purpose": "fixture",
            "paper_sections": ["results"], "process_summary": "test process", "human_review": "independently checked fixture"}
        log = {"compliance": {"ai_usage": [entry]}}
        pack = {"id": "cumcm", "rules": {"ai_declaration_position": {"value": "before_references"},
            "ai_details_filename": {"value": "AI工具使用详情.pdf"}}, "required_checks": ["ai_disclosure_passed"]}
        text = render_cumcm_use_statement(validate_entries([entry], "cumcm")) + "参考文献"
        report, _ = _ai_checks(self.root, log, pack, {"files": []}, text, [])
        self.assertFalse(report["passed"])
        self.assertTrue(any("actual supporting ZIP" in e for e in report["errors"]), report)


if __name__ == "__main__":
    unittest.main()
