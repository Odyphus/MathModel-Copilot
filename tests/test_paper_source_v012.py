"""Native source readback tests; fixtures are not actual human reviews."""
from __future__ import annotations
import copy
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))
import test_copilot_delivery as delivery_fixtures
from copilot_paper_source import audit_paper_source
from copilot_runtime import bind_file
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement


class PaperSourceV012Tests(unittest.TestCase):
    def setUp(self):
        self.fixture = delivery_fixtures.DeliveryTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.root, self.rt, self.delivery = self.fixture.root, self.fixture.rt, self.fixture.delivery
        self.section = self.fixture.section
        self.claim, self.text = self.fixture.claim, self.fixture.claim_text
        self.contract = {"version": "0.1", "sections": [self.section]}

    def section_text(self, text, **kwargs):
        self.fixture.write("current-section.md", text)
        self.section = self.delivery.section(self.fixture.rev(), "paper.results", "current-section.md", [self.claim], **kwargs)["result"]["section_id"]
        self.contract["sections"] = [self.section]

    def paragraph(self):
        return self.text + " [[claim:" + self.claim + "]]"

    def save(self, doc):
        path = self.root / "actual.docx"
        doc.save(path)
        return path

    def check(self, path=None, cp=None):
        return audit_paper_source(self.root, cp or self.rt.read()["copilot"], [self.section],
                                 path or self.root / "paper.docx", self.contract)

    def test_existing_registered_paragraph_matches_actual_docx(self):
        result = self.check()
        self.assertTrue(result["passed"], result)
        self.assertEqual(result["source_block_count"], 1)

    def test_extra_false_result_in_final_docx_is_rejected(self):
        doc = Document(self.root / "paper.docx")
        doc.add_paragraph("A fabricated result is 999 m.")
        self.assertFalse(self.check(self.save(doc))["passed"])

    def test_missing_registered_prose_is_rejected(self):
        self.section_text("Model scope is limited.\n\n" + self.paragraph())
        self.assertFalse(self.check()["passed"])

    def test_reordered_paragraphs_are_rejected(self):
        self.section_text("# 1 Model results\n\n" + self.paragraph())
        doc = Document()
        doc.add_paragraph(self.text)
        doc.add_heading("1 Model results", 1)
        self.assertFalse(self.check(self.save(doc))["passed"])

    def test_native_table_cells_are_compared(self):
        self.section_text("| Finding |\n| --- |\n| " + self.paragraph() + " |")
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "Finding"
        table.cell(1, 0).text = self.text
        path = self.save(doc)
        self.assertTrue(self.check(path)["passed"])
        table.cell(1, 0).text = "Computed length is 999 m."
        self.assertFalse(self.check(self.save(doc))["passed"])

    def test_native_math_tokens_and_fraction_order_are_compared(self):
        self.section_text(r"$$\frac{x}{y}$$" + "\n\n" + self.paragraph())
        doc = Document()
        paragraph = doc.add_paragraph()
        math = OxmlElement("m:oMath")
        fraction = OxmlElement("m:f")
        nodes = []
        for tag, value in (("m:num", "x"), ("m:den", "y")):
            part, run, text = OxmlElement(tag), OxmlElement("m:r"), OxmlElement("m:t")
            text.text = value
            run.append(text); part.append(run); fraction.append(part); nodes.append(text)
        math.append(fraction); paragraph._p.append(math)
        doc.add_paragraph(self.text)
        self.assertTrue(self.check(self.save(doc))["passed"])
        nodes[0].text, nodes[1].text = "y", "x"
        self.assertFalse(self.check(self.save(doc))["passed"])

    def test_bound_figure_bytes_are_compared(self):
        from PIL import Image
        path = self.root / "figure.png"
        Image.new("RGB", (24, 24), "navy").save(path)
        result_id = next(o["id"] for o in self.rt.read()["copilot"]["objects"].values() if o["kind"] == "ResultRecord")
        artifact = self.rt.register(self.fixture.rev(), "ArtifactRecord", "figure.test",
            {"artifact_id": "figure.test", "artifact_type": "figure", "path": "figure.png"},
            dependencies=[result_id], files=["figure.png"])["result"]["object_id"]
        self.section_text("[[figure:1]]\n\n![Source](figure.png)\n\n" + self.paragraph(),
                          structure=[{"kind": "figure", "number": 1, "artifact_id": artifact}])
        doc = Document()
        doc.add_paragraph("图 1")
        doc.add_picture(str(path))
        doc.add_paragraph(self.text)
        self.assertTrue(self.check(self.save(doc))["passed"])
        other = self.root / "unbound.png"
        Image.new("RGB", (24, 24), "red").save(other)
        doc = Document(); doc.add_paragraph("图 1"); doc.add_picture(str(other)); doc.add_paragraph(self.text)
        self.assertFalse(self.check(self.save(doc))["passed"])

    def test_self_reported_expected_text_cannot_authorize_changed_paper(self):
        self.contract["expected_text"] = "A fabricated result is 999 m."
        self.assertFalse(self.check()["passed"])

    def test_current_files_are_revalidated_despite_forged_status_in_memory(self):
        cp = copy.deepcopy(self.rt.read()["copilot"])
        path = cp["objects"][self.section]["payload"]["path"]
        self.fixture.write(path, self.paragraph() + " Fabricated result is 999 m.")
        cp["objects"][self.section]["files"] = [bind_file(self.root, path)]
        self.assertFalse(self.check(cp=cp)["passed"])

    def test_upstream_drift_is_rejected_even_if_docx_unchanged(self):
        solver = self.root / "solver.py"
        solver.write_text(solver.read_text(encoding="utf-8") + "\n# changed\n", encoding="utf-8")
        self.assertFalse(self.check()["passed"])

    def test_code_supplement_is_derived_from_actual_declared_source(self):
        self.contract["supplements"] = [{"kind": "source_code", "source_id": self.fixture.ids["code"], "path": "solver.py"}]
        doc = Document(self.root / "paper.docx")
        doc.styles.add_style("CopilotSourceCode", WD_STYLE_TYPE.PARAGRAPH)
        doc.add_paragraph("源程序：solver.py")
        code = doc.add_paragraph((self.root / "solver.py").read_text(encoding="utf-8").rstrip("\n"), "CopilotSourceCode")
        self.assertTrue(self.check(self.save(doc))["passed"])
        code.text += "\nprint('unbound code')"
        self.assertFalse(self.check(self.save(doc))["passed"])

    def test_unbound_code_file_cannot_be_a_supplement(self):
        self.fixture.write("new.py", "print(999)")
        self.contract["supplements"] = [{"kind": "source_code", "source_id": self.fixture.ids["code"], "path": "new.py"}]
        self.assertFalse(self.check()["passed"])

    def test_extra_contract_sections_cannot_disappear(self):
        self.contract["sections"].append("paper.missing@1")
        self.assertFalse(self.check()["passed"])

    def file_list(self):
        self.contract["supplements"] = [{"kind": "file_list", "source_id": self.fixture.ids["code"], "paths": ["solver.py"]}]
        doc = Document(self.root / "paper.docx")
        doc.add_paragraph("文件：solver.py")
        return doc

    def test_bound_file_list_is_derived_from_actual_source_files(self):
        self.assertTrue(self.check(self.save(self.file_list()))["passed"])

    def test_file_list_rejects_missing_final_path(self):
        self.file_list()
        self.assertFalse(self.check()["passed"])

    def test_file_list_rejects_unbound_or_unsafe_paths(self):
        self.file_list()
        self.fixture.write("unbound.py", "print(999)")
        for relative in ("unbound.py", "../solver.py", "C:/solver.py"):
            with self.subTest(path=relative):
                self.contract["supplements"][0]["paths"] = [relative]
                self.assertFalse(self.check()["passed"])

    def test_file_list_rejects_duplicate_or_empty_paths(self):
        self.file_list()
        for paths in (["solver.py", "solver.py"], []):
            with self.subTest(paths=paths):
                self.contract["supplements"][0]["paths"] = paths
                self.assertFalse(self.check()["passed"])

    def test_file_list_rechecks_physical_source_drift(self):
        doc = self.file_list()
        self.assertTrue(self.check(self.save(doc))["passed"])
        path = self.root / "solver.py"
        path.write_text(path.read_text(encoding="utf-8") + "\n# changed after registration\n", encoding="utf-8")
        self.assertFalse(self.check(self.save(doc))["passed"])

    def test_docx_replacement_during_readback_is_rejected(self):
        from unittest.mock import patch
        import copilot_paper_source
        original = copilot_paper_source.read_docx_blocks
        def replace_after_read(path):
            blocks = original(path)
            doc = Document(path)
            doc.add_paragraph("A fabricated additional result is 999 m.")
            doc.save(path)
            return blocks
        with patch.object(copilot_paper_source, "read_docx_blocks", replace_after_read):
            result = self.check()
        self.assertFalse(result["passed"])
        self.assertIn("DOCX changed during source readback", result["errors"])

    def hidden_style_document(self, kind=WD_STYLE_TYPE.PARAGRAPH):
        doc = Document(self.root / "paper.docx")
        style = doc.styles.add_style("HiddenSourceStyle", kind)
        style.font.hidden = True
        return doc, style

    def test_hidden_paragraph_style_cannot_establish_visible_coverage(self):
        doc, style = self.hidden_style_document()
        doc.paragraphs[0].style = style
        self.assertFalse(self.check(self.save(doc))["passed"])

    def test_hidden_character_style_cannot_establish_visible_coverage(self):
        doc, style = self.hidden_style_document(WD_STYLE_TYPE.CHARACTER)
        doc.paragraphs[0].runs[0].style = style
        self.assertFalse(self.check(self.save(doc))["passed"])

    def test_inherited_hidden_style_cannot_establish_visible_coverage(self):
        doc, parent = self.hidden_style_document()
        child = doc.styles.add_style("InheritedSourceStyle", WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = parent
        doc.paragraphs[0].style = child
        self.assertFalse(self.check(self.save(doc))["passed"])

    def test_hidden_document_default_cannot_establish_visible_coverage(self):
        from docx.oxml.ns import qn
        doc = Document(self.root / "paper.docx")
        defaults = doc.styles.element.find(qn("w:docDefaults"))
        properties = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
        properties.append(OxmlElement("w:vanish"))
        self.assertFalse(self.check(self.save(doc))["passed"])

    def test_explicit_visible_style_is_supported(self):
        doc, style = self.hidden_style_document()
        style.font.hidden = False
        doc.paragraphs[0].style = style
        self.assertTrue(self.check(self.save(doc))["passed"])

    def test_unused_hidden_style_is_conservatively_rejected(self):
        doc, _ = self.hidden_style_document()
        self.assertFalse(self.check(self.save(doc))["passed"])


if __name__ == "__main__":
    unittest.main()
